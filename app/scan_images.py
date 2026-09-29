"""Tool quét mã hàng Shopee: tải ảnh sản phẩm về máy, mỗi sản phẩm 1 thư mục.

Cách dùng (Windows: nháy đúp tai_anh_shopee.bat):
    python -m app.scan_images danh_sach_link.txt              # hoặc file .xlsx / .csv / .docx
    python -m app.scan_images danh_sach_link.txt --browser    # dùng trình duyệt khi Shopee chặn

Mỗi dòng / ô trong file có thể là:
    link sản phẩm   https://shopee.vn/Ten-san-pham-i.123456.7890123
                    https://shopee.vn/product/123456/7890123
    link rút gọn    https://s.shopee.vn/xxxx  (link aff cũng được)
    mã hàng         123456.7890123  hoặc  123456/7890123  (mã shop . mã sản phẩm)

Kết quả trong thư mục anh_san_pham/:
    <mã sản phẩm> - <tên sản phẩm>/01.jpg, 02.jpg ..., thong_tin.txt
    ket_qua.csv: danh sách sản phẩm, số ảnh tải được, lỗi (nếu có)
Chạy lại lần sau sẽ bỏ qua sản phẩm đã có ảnh (thêm --tai-lai để tải lại).
"""
import argparse
import csv
import json
import random
import re
import sys
import time
from io import BytesIO
from pathlib import Path

import httpx
from PIL import Image

from app.services import shopee

CODE_RE = re.compile(r"(?<![\d.])(\d{4,12})[./](\d{6,14})(?![\d.])")
SHORT_RE = re.compile(r"https?://(?:s\.shopee\.vn|shope\.ee|shp\.ee|vn\.shp\.ee)/[\w-]+", re.I)
LINK_RE = re.compile(r"https?://(?:www\.)?shopee\.vn/[^\s\"'<>]+", re.I)
_IMG_HEADERS = {"User-Agent": shopee._BROWSER_HEADERS["User-Agent"], "Referer": "https://shopee.vn/"}


def find_links(text: str) -> list[str]:
    """Tìm mọi link / mã hàng Shopee trong 1 đoạn chữ (1 dòng, 1 ô Excel...)."""
    found = []
    for m in SHORT_RE.finditer(text):
        found.append(m.group(0))
    for m in LINK_RE.finditer(text):
        link = m.group(0).rstrip(".,;)")
        if shopee.parse_ids(link):
            found.append(link)
    if not found:
        for m in CODE_RE.finditer(text):
            found.append(f"https://shopee.vn/product/{m.group(1)}/{m.group(2)}")
    return found


def read_input(path: Path) -> list[str]:
    """Đọc danh sách link từ file .txt / .csv / .xlsx / .docx, bỏ trùng, giữ thứ tự."""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        texts = [str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c]
    elif suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        texts = [p.text for p in doc.paragraphs]
        texts += [cell.text for t in doc.tables for row in t.rows for cell in row.cells]
    else:
        raw = path.read_bytes()
        try:
            texts = raw.decode("utf-8-sig").splitlines()
        except UnicodeDecodeError:
            texts = raw.decode("cp1258", errors="ignore").splitlines()
    links = [link for t in texts for link in find_links(t)]
    return list(dict.fromkeys(links))


def safe_name(text: str, limit: int = 60) -> str:
    """Tên thư mục hợp lệ trên Windows (giữ tiếng Việt có dấu)."""
    text = re.sub(r'[\\/:*?"<>|\r\n\t]+', " ", text or "")
    text = re.sub(r"\s+", " ", text).strip()[:limit]
    return text.rstrip(". ") or "san-pham"


def _collect_http(link: str) -> dict:
    """Lấy tên, giá, ảnh... bằng cách gọi trực tiếp (nhanh, nhưng Shopee hay chặn)."""
    info = {"input": link, "images": []}
    product_link = shopee.resolve_link(link)
    info["product_link"] = product_link
    ids = shopee.parse_ids(product_link)
    if not ids:
        info["error"] = "Không đọc được mã sản phẩm (link rút gọn không mở được?)"
        return info
    info["shop_id"], info["item_id"] = ids
    item = shopee._public_item(ids[0], ids[1], product_link)
    if item:
        _fill_from_item(info, item)
    if len(info["images"]) < 2:
        api = shopee.lookup_product(product_link)            # chỉ có khi đã điền SHOPEE_APP_ID/SECRET
        if api.get("image_url"):
            info["images"].append(api["image_url"])
            info.setdefault("name", api.get("name"))
        og = shopee._og_meta(product_link)
        if og.get("og:image"):
            info["images"].append(og["og:image"])
        if og.get("og:title") and not info.get("name"):
            info["name"] = re.sub(r"\s*\|\s*Shopee.*$", "", og["og:title"]).strip()
    info.setdefault("name", shopee.name_from_link(product_link))
    info["images"] = list(dict.fromkeys(info["images"]))
    return info


def _fill_from_item(info: dict, item: dict) -> None:
    """Điền thông tin từ dữ liệu sản phẩm của Shopee (dạng api/v4/item/get hoặc pdp/get_pc)."""
    info["name"] = item.get("name") or info.get("name")
    price = item.get("price") or item.get("price_min")
    if isinstance(price, (int, float)) and price:
        info["price"] = int(price / 100000)                   # Shopee lưu giá x100000
    info["sold"] = item.get("historical_sold") or item.get("sold") or info.get("sold")
    rating = (item.get("item_rating") or {}).get("rating_star")
    if rating:
        info["rating"] = round(float(rating), 1)
    info["description"] = item.get("description") or info.get("description", "")
    images = [shopee.IMG_CDN + h for h in item.get("images") or [] if isinstance(h, str)]
    for model in item.get("tier_variations") or []:           # ảnh từng phân loại (màu, mẫu...)
        images += [shopee.IMG_CDN + h for h in model.get("images") or [] if isinstance(h, str) and h]
    info["images"] = list(dict.fromkeys(info.get("images", []) + images))


def _find_item(data, item_id: str):
    """Tìm khối dữ liệu sản phẩm (có "images" và đúng itemid) trong JSON Shopee trả về."""
    if isinstance(data, dict):
        if isinstance(data.get("images"), list) and str(data.get("itemid") or data.get("item_id")) == item_id:
            return data
        for v in data.values():
            found = _find_item(v, item_id)
            if found:
                return found
    elif isinstance(data, list):
        for v in data:
            found = _find_item(v, item_id)
            if found:
                return found
    return None


class BrowserCollector:
    """Mở trang sản phẩm bằng Edge / Chrome có sẵn trên máy (như người dùng thật) và đọc dữ liệu.

    Dùng hồ sơ trình duyệt riêng (thư mục data/shopee_profile): lần đầu nếu Shopee đòi đăng nhập
    thì bạn đăng nhập ngay trong cửa sổ đó, các lần sau không phải đăng nhập lại.
    """

    def __init__(self, profile_dir: Path):
        try:
            from playwright.sync_api import sync_playwright
        except ImportError:
            raise SystemExit("Chưa cài thư viện trình duyệt. Chạy: pip install playwright  (rồi chạy lại)")
        self._pw = sync_playwright().start()
        last_error = None
        for channel in ("msedge", "chrome", None):
            try:
                self.ctx = self._pw.chromium.launch_persistent_context(
                    str(profile_dir), channel=channel, headless=False, locale="vi-VN",
                    viewport={"width": 1280, "height": 900})
                break
            except Exception as e:  # noqa: BLE001 - máy không có Edge thì thử Chrome
                last_error = e
        else:
            raise SystemExit(f"Không mở được Edge / Chrome: {last_error}")
        self.page = self.ctx.pages[0] if self.ctx.pages else self.ctx.new_page()

    def collect(self, link: str, wait_login: bool = True) -> dict:
        info = {"input": link, "images": []}
        captured = []

        def on_response(resp):
            if "/api/v4/pdp/get" in resp.url or "/api/v4/item/get" in resp.url:
                try:
                    captured.append(resp.json())
                except Exception:  # noqa: BLE001
                    pass

        self.page.on("response", on_response)
        try:
            self.page.goto(link, wait_until="domcontentloaded", timeout=60000)
            self.page.wait_for_timeout(4000)
            if wait_login and "/buyer/login" in self.page.url:
                print("   Shopee đòi đăng nhập: hãy đăng nhập trong cửa sổ trình duyệt vừa mở"
                      " (chỉ cần 1 lần), xong quay lại đây bấm Enter...")
                input()
                captured.clear()
                self.page.goto(link, wait_until="domcontentloaded", timeout=60000)
                self.page.wait_for_timeout(4000)
            product_link = self.page.url.split("?")[0]
            info["product_link"] = product_link
            ids = shopee.parse_ids(product_link) or shopee.parse_ids(link)
            if not ids:
                info["error"] = "Trang mở ra không phải trang sản phẩm Shopee"
                return info
            info["shop_id"], info["item_id"] = ids
            for data in captured:
                item = _find_item(data, ids[1])
                if item:
                    _fill_from_item(info, item)
                    break
            if len(info["images"]) < 2:                         # đọc thẳng ảnh trên trang
                srcs = self.page.eval_on_selector_all(
                    "img", "els => els.map(e => e.currentSrc || e.src)")
                hashes = [re.search(r"/file/([\w-]+)", s) for s in srcs if "susercontent.com/file/" in s]
                info["images"] += [shopee.IMG_CDN + m.group(1) for m in hashes if m][:10]
                og = self.page.eval_on_selector('meta[property="og:image"]', "e => e.content") \
                    if self.page.query_selector('meta[property="og:image"]') else ""
                if og:
                    info["images"].insert(0, og)
                info["images"] = list(dict.fromkeys(info["images"]))
            if not info.get("name"):
                title = self.page.title()
                info["name"] = re.sub(r"\s*\|\s*Shopee.*$", "", title).strip()
        except Exception as e:  # noqa: BLE001
            info["error"] = f"Lỗi mở trang: {e}"
        finally:
            self.page.remove_listener("response", on_response)
        return info

    def close(self):
        self.ctx.close()
        self._pw.stop()


def download(url: str, dest: Path) -> bool:
    """Tải 1 ảnh, kiểm tra đúng là ảnh, lưu JPEG chất lượng cao."""
    try:
        with httpx.Client(timeout=30, headers=_IMG_HEADERS, follow_redirects=True) as client:
            resp = client.get(url)
        resp.raise_for_status()
        im = Image.open(BytesIO(resp.content))
        im.load()
        if min(im.size) < 200:                                  # bỏ icon / ảnh quá nhỏ
            return False
        im.convert("RGB").save(dest, "JPEG", quality=95)
        return True
    except Exception:  # noqa: BLE001
        return False


def _write_info(folder: Path, info: dict) -> None:
    lines = [
        f"Tên: {info.get('name') or ''}",
        f"Link sản phẩm: {info.get('product_link') or ''}",
        f"Link đưa vào: {info.get('input') or ''}",
        f"Mã shop: {info.get('shop_id') or ''}   Mã sản phẩm: {info.get('item_id') or ''}",
    ]
    if info.get("price"):
        lines.append(f"Giá: {info['price']:,}đ".replace(",", "."))
    if info.get("sold"):
        lines.append(f"Đã bán: {info['sold']}")
    if info.get("rating"):
        lines.append(f"Đánh giá: {info['rating']}/5")
    if info.get("description"):
        lines += ["", "Mô tả:", info["description"]]
    (folder / "thong_tin.txt").write_text("\n".join(lines), encoding="utf-8")


def _existing_folder(out_dir: Path, item_id: str) -> Path | None:
    for folder in out_dir.glob(f"{item_id} - *"):
        if any(folder.glob("*.jpg")):
            return folder
    return None


def scan(links: list[str], out_dir: Path, max_images: int = 10, use_browser: bool = False,
         redownload: bool = False, delay: tuple[float, float] = (1.5, 4.0), log=print) -> list[dict]:
    """Quét danh sách link, tải ảnh. Trả về danh sách kết quả từng sản phẩm."""
    out_dir.mkdir(parents=True, exist_ok=True)
    browser = None
    results = []
    try:
        for i, link in enumerate(links, 1):
            log(f"[{i}/{len(links)}] {link}")
            ids = shopee.parse_ids(link)
            if ids and not redownload:
                folder = _existing_folder(out_dir, ids[1])
                if folder:
                    n = len(list(folder.glob("*.jpg")))
                    log(f"   Đã có {n} ảnh, bỏ qua ({folder.name})")
                    results.append({"input": link, "item_id": ids[1], "name": folder.name.split(" - ", 1)[-1],
                                    "folder": str(folder), "downloaded": n, "status": "đã có"})
                    continue
            info = {} if use_browser else _collect_http(link)
            if use_browser or len(info.get("images", [])) < 2:
                if browser is None and use_browser:
                    browser = BrowserCollector(out_dir.parent / "data" / "shopee_profile")
                if browser:
                    got = browser.collect(link)
                    if len(got.get("images", [])) >= len(info.get("images", [])):
                        info = got
            result = {"input": link, "item_id": info.get("item_id", ""), "name": info.get("name", ""),
                      "product_link": info.get("product_link", ""), "found": len(info.get("images", [])),
                      "downloaded": 0, "folder": "", "status": ""}
            if not info.get("item_id"):
                result["status"] = info.get("error") or "Không đọc được sản phẩm"
                log(f"   LỖI: {result['status']}")
                results.append(result)
                continue
            folder = out_dir / f"{info['item_id']} - {safe_name(info.get('name') or '')}"
            folder.mkdir(parents=True, exist_ok=True)
            _write_info(folder, info)
            n = 0
            for url in info["images"][:max_images]:
                if download(url, folder / f"{n + 1:02d}.jpg"):
                    n += 1
            result.update(downloaded=n, folder=str(folder))
            if n:
                result["status"] = "ok" if n >= 2 else "chỉ lấy được 1 ảnh (thử lại với --browser)"
            else:
                result["status"] = info.get("error") or "Shopee chặn, chưa lấy được ảnh (thử lại với --browser)"
            log(f"   {info.get('name') or ''}: tải {n} ảnh -> {folder.name}")
            results.append(result)
            if i < len(links):
                time.sleep(random.uniform(*delay))              # nghỉ giữa các sản phẩm để Shopee không chặn
    finally:
        if browser:
            browser.close()
    with (out_dir / "ket_qua.csv").open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["Link đưa vào", "Mã sản phẩm", "Tên", "Số ảnh tải được", "Thư mục", "Trạng thái"])
        for r in results:
            writer.writerow([r["input"], r["item_id"], r["name"], r["downloaded"], r["folder"], r["status"]])
    return results


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Tải ảnh sản phẩm Shopee về máy, mỗi sản phẩm 1 thư mục.")
    parser.add_argument("file", nargs="?", default="danh_sach_link.txt",
                        help="file chứa link / mã hàng (.txt, .csv, .xlsx, .docx)")
    parser.add_argument("--out", default="anh_san_pham", help="thư mục lưu ảnh (mặc định: anh_san_pham)")
    parser.add_argument("--max", type=int, default=10, help="số ảnh tối đa mỗi sản phẩm (mặc định 10)")
    parser.add_argument("--browser", action="store_true", help="mở Edge/Chrome để lấy dữ liệu khi Shopee chặn")
    parser.add_argument("--tai-lai", action="store_true", help="tải lại cả sản phẩm đã có ảnh")
    args = parser.parse_args(argv)

    path = Path(args.file)
    if not path.exists():
        path.write_text("# Dán link sản phẩm Shopee hoặc mã hàng (mã shop.mã sản phẩm), mỗi dòng 1 cái\n",
                        encoding="utf-8")
        print(f"Đã tạo file {path}. Mở file, dán link Shopee vào (mỗi dòng 1 link), lưu lại rồi chạy lại tool.")
        return 1
    links = read_input(path)
    if not links:
        print(f"Không tìm thấy link / mã hàng Shopee nào trong {path}.")
        return 1
    print(f"Tìm thấy {len(links)} sản phẩm trong {path}. Bắt đầu tải ảnh...\n")
    results = scan(links, Path(args.out), args.max, args.browser, args.tai_lai)
    ok = sum(1 for r in results if r["downloaded"])
    total = sum(r["downloaded"] for r in results)
    print(f"\nXong: {ok}/{len(results)} sản phẩm có ảnh, tổng {total} ảnh. Xem thư mục {args.out}/"
          f" và file {args.out}/ket_qua.csv")
    blocked = [r for r in results if r["downloaded"] < 2 and r["item_id"]]
    if blocked and not args.browser:
        print(f"{len(blocked)} sản phẩm lấy được ít ảnh (Shopee chặn). Chạy lại bằng chế độ trình duyệt:"
              f" tai_anh_shopee.bat --browser  (hoặc thêm --browser)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
