"""Tool săn sản phẩm aff Shopee hằng ngày: hoa hồng cao + đang bán chạy -> lấy link aff về file Excel.

Cần Shopee Affiliate Open API (SHOPEE_APP_ID, SHOPEE_APP_SECRET trong file .env).
Đăng ký tại trang Shopee Affiliate (affiliate.shopee.vn) -> mục Open API.

Cách dùng (Windows: nháy đúp san_aff_shopee.bat):
    python -m app.aff_hunter                     # đọc tu_khoa.txt, ghi san_pham_aff/<ngày>.xlsx
    python -m app.aff_hunter --top 20            # số sản phẩm mỗi nhóm
    python -m app.aff_hunter --demo              # chạy thử bằng dữ liệu giả (chưa có Open API)

File tu_khoa.txt: dòng [Tên nhóm] rồi các từ khoá, mỗi dòng 1 từ khoá:
    [Thú cưng]
    hạt cho mèo
    cát vệ sinh mèo

Mỗi từ khoá lấy 2 danh sách từ Shopee: hoa hồng cao nhất và bán chạy nhất, gộp lại, lọc theo
% hoa hồng / lượt bán / số sao tối thiểu, rồi xếp hạng theo "tiền hoa hồng kỳ vọng":
giá x % hoa hồng, nhân hệ số tin cậy từ lượt bán và đánh giá.
"""
import argparse
import re
import sys
from datetime import datetime
from pathlib import Path

from app import config
from app.services import shopee

SUB_ID = "hunt"


def read_keywords(path: Path) -> list[tuple[str, str]]:
    """[(nhóm, từ khoá)] theo thứ tự trong file."""
    group, out = "Chung", []
    for raw in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        m = re.fullmatch(r"\[(.+)\]", line)
        if m:
            group = m.group(1).strip()
        else:
            out.append((group, line))
    return list(dict.fromkeys(out))


def sample_keywords() -> str:
    from app.db import DEFAULT_NICHES

    lines = ["# Dòng [Tên nhóm] rồi các từ khoá tìm trên Shopee, mỗi dòng 1 từ khoá. Dòng # là ghi chú."]
    for niche, words in DEFAULT_NICHES.items():
        lines += [f"[{niche}]", *words, ""]
    return "\n".join(lines)


def expected_commission(p: dict) -> float:
    """Tiền hoa hồng kỳ vọng mỗi đơn (đồng), nhân hệ số tin cậy từ lượt bán & đánh giá."""
    per_order = p["price"] * p["commission_rate"]
    trust = min(p["sales"], 20000) ** 0.5 / 141 * (p["rating"] / 5)
    return round(per_order * (0.5 + trust), 1)


def hunt(keywords: list[tuple[str, str]], min_rate: float = 0.08, min_sales: int = 500,
         min_rating: float = 4.7, top: int = 15, per_keyword: int = 50, blacklist: list[str] | None = None,
         log=print) -> list[dict]:
    """Tìm, lọc, xếp hạng sản phẩm. Mỗi nhóm giữ `top` sản phẩm điểm cao nhất."""
    blacklist = [w.lower() for w in blacklist or []]
    by_group: dict[str, dict[str, dict]] = {}
    for i, (group, kw) in enumerate(keywords, 1):
        log(f"[{i}/{len(keywords)}] {group}: {kw}")
        found = []
        for sort in (shopee.SORT_COMMISSION_DESC, shopee.SORT_SALES_DESC):
            try:
                found += shopee.search_products(kw, per_keyword, sort)
            except Exception as e:  # noqa: BLE001 - 1 từ khoá lỗi không làm dừng cả lượt
                log(f"   LỖI: {e}")
        kept = 0
        for p in found:
            if (p["commission_rate"] < min_rate or p["sales"] < min_sales or p["rating"] < min_rating
                    or any(w in p["name"].lower() for w in blacklist)):
                continue
            bucket = by_group.setdefault(group, {})
            if p["item_id"] not in bucket:
                bucket[p["item_id"]] = {**p, "group": group, "keyword": kw, "score": expected_commission(p)}
                kept += 1
        log(f"   {len(found)} kết quả, {kept} sản phẩm đạt chuẩn")
    result = []
    for group, items in by_group.items():
        result += sorted(items.values(), key=lambda p: p["score"], reverse=True)[:top]
    return result


def aff_link(product_link: str, item_id: str, day: str) -> str:
    """Link aff rút gọn, gắn subId = ngày săn để biết đơn đến từ danh sách ngày nào."""
    sub_ids = [SUB_ID, day.replace("-", "")]
    if not config.SHOPEE_ENABLED:
        return f"https://s.shopee.vn/demo{item_id}?sub={'-'.join(sub_ids)}"
    data = shopee._call(shopee.LINK_MUTATION, {"url": product_link, "subIds": sub_ids})
    return data["generateShortLink"]["shortLink"]


def add_aff_links(products: list[dict], day: str, log=print) -> None:
    for p in products:
        link = p.get("product_link") or ""
        try:
            p["aff_link"] = aff_link(link, p["item_id"], day) if link else ""
        except Exception as e:  # noqa: BLE001 - lỗi tạo link thì dùng link offer Shopee trả về
            log(f"   Không tạo được link aff cho {p['name'][:40]}: {e}")
            p["aff_link"] = p.get("offer_link") or ""


def load_seen(path: Path) -> set[str]:
    return set(path.read_text(encoding="utf-8").split()) if path.exists() else set()


def write_excel(path: Path, products: list[dict], seen: set[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "San pham aff"
    head = ["Nhóm", "Mới?", "Tên sản phẩm", "Giá (đ)", "% hoa hồng", "Hoa hồng/đơn (đ)", "Đã bán", "Sao",
            "Shop", "Link aff", "Link sản phẩm", "Ảnh", "Từ khoá"]
    ws.append(head)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="EE4D2D")
    for p in products:
        ws.append([p["group"], "MỚI" if p["item_id"] not in seen else "", p["name"], int(p["price"]),
                   round(p["commission_rate"] * 100, 1), int(p["price"] * p["commission_rate"]), p["sales"],
                   p["rating"], p["shop_name"], p.get("aff_link", ""), p["product_link"], p["image_url"],
                   p["keyword"]])
    for col, width in zip("ABCDEFGHIJKLM", (16, 7, 50, 11, 11, 15, 9, 6, 22, 34, 40, 30, 18)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def write_links(path: Path, products: list[dict]) -> None:
    """File link aff theo nhóm, dán thẳng vào app / tool tải ảnh được."""
    lines, group = [], None
    for p in products:
        if p["group"] != group:
            group = p["group"]
            lines += ["", f"[{group}]"] if lines else [f"[{group}]"]
        lines.append(p.get("aff_link") or p["product_link"])
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Săn sản phẩm Shopee hoa hồng cao, bán chạy và lấy link aff.")
    ap.add_argument("file", nargs="?", default="tu_khoa.txt", help="file từ khoá (mặc định tu_khoa.txt)")
    ap.add_argument("--out", default="san_pham_aff", help="thư mục lưu kết quả (mặc định san_pham_aff)")
    ap.add_argument("--top", type=int, default=15, help="số sản phẩm tốt nhất mỗi nhóm (mặc định 15)")
    ap.add_argument("--hoa-hong", type=float, default=8, help="%% hoa hồng tối thiểu (mặc định 8)")
    ap.add_argument("--da-ban", type=int, default=500, help="lượt bán tối thiểu (mặc định 500)")
    ap.add_argument("--sao", type=float, default=4.7, help="số sao tối thiểu (mặc định 4.7)")
    ap.add_argument("--demo", action="store_true", help="chạy thử bằng dữ liệu giả, không cần Open API")
    ap.add_argument("--tu-dong", action="store_true", help="chạy theo lịch hằng ngày (không khác chạy tay)")
    args = ap.parse_args(argv)

    if not config.SHOPEE_ENABLED and not args.demo:
        print("Chưa có Shopee Affiliate Open API.\n"
              "1. Vào https://affiliate.shopee.vn , đăng nhập tài khoản affiliate, tìm mục Open API và đăng ký.\n"
              "2. Được duyệt sẽ có AppID và Secret: điền vào file .env:\n"
              "     SHOPEE_APP_ID=...\n     SHOPEE_APP_SECRET=...\n"
              "3. Chạy lại tool. Muốn xem thử trước: thêm --demo (dữ liệu giả).")
        return 1
    path = Path(args.file)
    if not path.exists():
        path.write_text(sample_keywords(), encoding="utf-8")
        print(f"Đã tạo file {path} với từ khoá mẫu theo ngành hàng. Sửa lại cho đúng ngành bạn chạy rồi chạy lại.")
        return 1
    keywords = read_keywords(path)
    if not keywords:
        print(f"Chưa có từ khoá nào trong {path}.")
        return 1

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    day = datetime.now(config.TZ).strftime("%Y-%m-%d")
    seen_file = out / "da_thay.txt"
    seen = load_seen(seen_file)
    from app.db import DEFAULT_SETTINGS

    print(f"{'[DEMO - dữ liệu giả] ' if args.demo and not config.SHOPEE_ENABLED else ''}"
          f"Săn sản phẩm cho {len(keywords)} từ khoá...\n")
    products = hunt(keywords, args.hoa_hong / 100, args.da_ban, args.sao, args.top,
                    blacklist=DEFAULT_SETTINGS.get("blacklist"))
    add_aff_links(products, day)
    xlsx, txt = out / f"{day}.xlsx", out / f"{day}_link_aff.txt"
    write_excel(xlsx, products, seen)
    write_links(txt, products)
    new = [p for p in products if p["item_id"] not in seen]
    with seen_file.open("a", encoding="utf-8") as fh:
        fh.writelines(f"{p['item_id']}\n" for p in new)
    print(f"\nXong: {len(products)} sản phẩm ({len(new)} sản phẩm mới so với các ngày trước).\n"
          f"  Bảng chi tiết: {xlsx}\n  Danh sách link aff: {txt}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
