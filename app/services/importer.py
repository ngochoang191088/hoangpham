"""Đọc danh sách sản phẩm + link aff + video từ file Excel (.xlsx), CSV, Word (.docx) hoặc text.

Mỗi ngành hàng dùng 1 file riêng. Cột được nhận diện theo tên (không phân biệt hoa thường, có dấu hay không):
    Tên sản phẩm | Link sản phẩm | Link aff | Giá | Mô tả | Link ảnh | Link video

Nếu file không có dòng tiêu đề, app tự nhận diện theo nội dung từng ô:
link shopee.vn/... là link sản phẩm, s.shopee.vn / shope.ee là link aff, link .mp4 / Google Drive là video.
"""
import csv
import hashlib
import io
import re
import unicodedata

FIELDS = ["name", "product_link", "aff_link", "price", "description", "image_url", "video_url", "niche"]
FIELD_LABELS = {
    "name": "Tên sản phẩm", "product_link": "Link sản phẩm", "aff_link": "Link aff", "price": "Giá",
    "description": "Mô tả / điểm nổi bật", "image_url": "Link ảnh", "video_url": "Link video",
}

# Từ khoá trong tiêu đề cột (đã bỏ dấu) -> trường. Kiểm tra theo thứ tự.
_HEADER_RULES = [
    ("aff_link", ["aff", "affiliate", "tiep thi", "link rut gon", "short"]),
    ("video_url", ["video", "clip"]),
    ("image_url", ["anh", "hinh", "image", "img", "photo"]),
    ("product_link", ["link san pham", "link sp", "url san pham", "product link", "link goc", "link shopee",
                      "link", "url"]),
    ("price", ["gia", "price"]),
    ("description", ["mo ta", "diem noi bat", "noi dung", "description", "ghi chu", "thong tin", "uu diem"]),
    ("niche", ["nganh"]),
    ("name", ["ten", "san pham", "name", "tieu de", "title"]),
]

URL_RE = re.compile(r"https?://\S+")


def strip_accents(text: str) -> str:
    text = unicodedata.normalize("NFD", str(text or "")).replace("đ", "d").replace("Đ", "D")
    return "".join(c for c in text if unicodedata.category(c) != "Mn").lower().strip()


def _field_for_header(header: str) -> str | None:
    h = strip_accents(header)
    if not h:
        return None
    for field, keys in _HEADER_RULES:
        if any(k in h for k in keys):
            return field
    return None


def classify_url(url: str) -> str | None:
    u = url.lower()
    if re.search(r"(s\.shopee\.vn|shope\.ee|shp\.ee|shopee\.vn/universal-link|/an_redir|affiliate)", u):
        return "aff_link"
    if re.search(r"\.(mp4|mov|m4v|webm)(\?|$)", u) or "drive.google.com" in u:
        return "video_url"
    if re.search(r"\.(jpe?g|png|webp)(\?|$)", u) or "cf.shopee" in u or "susercontent" in u:
        return "image_url"
    if "shopee.vn" in u or "lazada" in u or "tiktok.com/view/product" in u:
        return "product_link"
    return None


def parse_price(value) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = strip_accents(value).replace("vnd", "").replace("d", "").strip()
    mult = 1000 if text.endswith("k") else 1
    digits = re.sub(r"[^\d]", "", text)
    return float(digits) * mult if digits else 0.0


def _rows_from_table(table: list[list]) -> list[dict]:
    """Bảng (dòng x cột) -> danh sách dict theo FIELDS."""
    table = [[("" if c is None else c) for c in row] for row in table if any(str(c or "").strip() for c in row)]
    if not table:
        return []
    header_map = {i: _field_for_header(str(h)) for i, h in enumerate(table[0])}
    has_header = sum(1 for f in header_map.values() if f) >= 2 and not any(
        URL_RE.search(str(h)) for h in table[0])
    rows = []
    for raw in (table[1:] if has_header else table):
        row: dict = {}
        for i, cell in enumerate(raw):
            text = str(cell).strip()
            if not text:
                continue
            field = header_map.get(i) if has_header else None
            if not has_header or field is None:
                # không có tiêu đề: đoán theo nội dung ô
                if URL_RE.fullmatch(text):
                    field = classify_url(text) or ("product_link" if "product_link" not in row else None)
                elif "name" not in row:
                    field = "name"
                elif parse_price(text) and "price" not in row and len(text) < 20:
                    field = "price"
                else:
                    field = "description"
            if field == "description" and row.get("description"):
                row["description"] += "\n" + text
            elif field and field not in row:
                row[field] = cell if field == "price" else text
        if row:
            rows.append(row)
    return rows


def _lines_to_rows(text: str) -> list[dict]:
    """Văn bản tự do: mỗi nhóm dòng (cách nhau dòng trống) là 1 sản phẩm; hoặc mỗi dòng có dấu |."""
    rows = []
    if "|" in text or "\t" in text:
        table = [re.split(r"\s*[|\t]\s*", line) for line in text.splitlines() if line.strip()]
        return _rows_from_table(table)
    for block in re.split(r"\n\s*\n", text):
        cells = [c.strip() for c in block.splitlines() if c.strip()]
        if cells:
            rows += _rows_from_table([cells])
    return rows


def read_file(filename: str, data: bytes) -> list[dict]:
    """Đọc file tải lên, trả về danh sách dòng sản phẩm (chưa kiểm tra)."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = []
        for ws in wb.worksheets:
            rows += _rows_from_table([list(r) for r in ws.iter_rows(values_only=True)])
        return rows
    if name.endswith(".docx"):
        from docx import Document

        doc = Document(io.BytesIO(data))
        rows = []
        for table in doc.tables:
            rows += _rows_from_table([[c.text for c in r.cells] for r in table.rows])
        if not rows:
            rows = _lines_to_rows("\n".join(p.text for p in doc.paragraphs))
        return rows
    text = data.decode("utf-8-sig", errors="replace")
    if name.endswith(".csv"):
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t") if text.strip() else csv.excel
        return _rows_from_table(list(csv.reader(io.StringIO(text), dialect)))
    if name.endswith(".xls"):
        raise ValueError("File .xls đời cũ chưa hỗ trợ, hãy lưu lại thành .xlsx")
    return _lines_to_rows(text)


def item_id_from_link(link: str) -> str | None:
    """Mã sản phẩm Shopee từ link (dạng ...-i.<shop>.<item> hoặc /product/<shop>/<item>)."""
    m = re.search(r"-i\.(\d+)\.(\d+)", link) or re.search(r"/product/(\d+)/(\d+)", link)
    return f"{m.group(1)}.{m.group(2)}" if m else None


def make_item_id(row: dict) -> str:
    link = row.get("product_link") or row.get("aff_link") or row.get("name", "")
    return item_id_from_link(link) or "u" + hashlib.md5(link.strip().lower().encode()).hexdigest()[:12]


def validate(row: dict) -> str | None:
    """Trả về lỗi nếu dòng thiếu dữ liệu bắt buộc."""
    if not (row.get("product_link") or row.get("aff_link")):
        return "thiếu link sản phẩm / link aff"
    if not row.get("aff_link"):
        return "thiếu link aff"
    return None


def template_xlsx(niche: str) -> bytes:
    """File Excel mẫu cho 1 ngành hàng."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = (niche or "San pham")[:31]
    headers = [FIELD_LABELS[f] for f in ["name", "product_link", "aff_link", "price", "description",
                                         "image_url", "video_url"]]
    ws.append(headers)
    ws.append(["Nồi chiên không dầu 5L", "https://shopee.vn/Noi-chien-khong-dau-5L-i.123456.7890123",
               "https://s.shopee.vn/AbCdEf", 899000, "Dung tích 5L, hẹn giờ 60 phút, lòng chống dính",
               "https://down-vn.img.susercontent.com/file/anh-san-pham.jpg",
               "https://drive.google.com/file/d/ID_VIDEO/view"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="EE4D2D")
    for col, width in zip("ABCDEFG", [30, 45, 28, 12, 45, 40, 45]):
        ws.column_dimensions[col].width = width
    note = wb.create_sheet("Huong dan")
    for line in [
        "Mỗi dòng là 1 sản phẩm. Bắt buộc: Link aff. Nên có: Tên sản phẩm, Link sản phẩm, Giá, Mô tả, Link ảnh.",
        "Link video (không bắt buộc): link .mp4 trực tiếp hoặc Google Drive (chia sẻ: Bất kỳ ai có đường liên kết).",
        "Có video thì app đăng video; không có video thì đăng ảnh sản phẩm.",
        "Một sản phẩm có nhiều video: thêm nhiều dòng cùng Link aff, mỗi dòng 1 link video.",
        "Chỉ dùng video/ảnh bạn có quyền sử dụng.",
    ]:
        note.append([line])
    note.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
