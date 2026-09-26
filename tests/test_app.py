import json
import os
import tempfile
from pathlib import Path

os.environ["DB_PATH"] = os.path.join(tempfile.mkdtemp(), "test.db")
os.environ["ADMIN_USER"] = "admin"
os.environ["ADMIN_PASSWORD"] = "secret"
os.environ["UPLOAD_DIR"] = os.path.join(tempfile.mkdtemp(), "uploads")
os.environ["MEDIA_DIR"] = os.path.join(tempfile.mkdtemp(), "media")
os.environ["SHOPEE_FETCH"] = "0"      # test không gọi mạng Shopee
for key in ("FB_APP_ID", "FB_APP_SECRET", "FB_SYSTEM_USER_TOKEN", "SHOPEE_APP_ID", "SHOPEE_APP_SECRET",
            "ANTHROPIC_API_KEY"):
    os.environ[key] = ""

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app import db, demo  # noqa: E402
from app.main import app  # noqa: E402
from app.services import ai_writer, facebook, importer, pipeline, shopee  # noqa: E402



@pytest.fixture(scope="module")
def client():
    demo.seed(n_pages=15, days=5, kit_niches=0)
    with TestClient(app) as c:
        r = c.post("/login", data={"password": "secret"})
        assert r.status_code == 200
        yield c


def test_requires_login():
    with TestClient(app) as anon:
        r = anon.get("/", follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] == "/login"
        assert anon.get("/login").status_code == 200
        r = anon.post("/login", data={"password": "wrong"}, follow_redirects=False)
        assert "Sai" in r.headers["location"] or "Sai" in anon.get(r.headers["location"]).text
        assert anon.get("/", follow_redirects=False).status_code == 303


@pytest.mark.parametrize("url", [
    "/", "/?days=7", "/pages", "/pages?sort=orders&niche=Gia%20d%E1%BB%A5ng", "/pages/100000000000",
    "/pages/100000000000?status_=pending", "/posts", "/posts?status=failed&flagged=1", "/review",
    "/review?flagged=1", "/products", "/settings",
])
def test_pages_render(client, url):
    assert client.get(url).status_code == 200


def test_unknown_page_404(client):
    assert client.get("/pages/nope").status_code == 404


def test_approve_then_publish(client):
    with db.get_conn() as conn:
        post = conn.execute(
            "SELECT posts.id FROM posts JOIN pages ON pages.id = posts.page_id "
            "WHERE posts.status='pending' AND pages.status='active' LIMIT 1").fetchone()
        conn.execute("UPDATE posts SET scheduled_at = ? WHERE id = ?", (db.now_iso(), post["id"]))
    client.post(f"/posts/{post['id']}/approve")
    with db.get_conn() as conn:
        assert conn.execute("SELECT status FROM posts WHERE id=?", (post["id"],)).fetchone()[0] == "approved"
        pipeline.publish_due(conn)
        row = conn.execute("SELECT status, fb_post_id FROM posts WHERE id=?", (post["id"],)).fetchone()
    assert row["status"] == "published" and row["fb_post_id"]


def test_global_pause_blocks_publishing(client):
    client.post("/pause")
    with db.get_conn() as conn:
        conn.execute("UPDATE posts SET status='approved', scheduled_at=? WHERE status='pending'", (db.now_iso(),))
        assert pipeline.publish_due(conn) == (0, 0)
    client.post("/pause")
    with db.get_conn() as conn:
        assert db.get_settings(conn)["global_pause"] is False


def test_bulk_reject(client):
    with db.get_conn() as conn:
        conn.execute("UPDATE posts SET status='pending' WHERE status='approved'")
        ids = [r[0] for r in conn.execute("SELECT id FROM posts WHERE status='pending' LIMIT 3")]
    client.post("/review/bulk", data={"action": "reject", "ids": ids})
    with db.get_conn() as conn:
        marks = ",".join("?" * len(ids))
        statuses = {r[0] for r in conn.execute(f"SELECT status FROM posts WHERE id IN ({marks})", ids)}
    assert statuses == {"rejected"}


def test_check_content_flags():
    s = dict(db.DEFAULT_SETTINGS)
    assert ai_writer.check_content("Sản phẩm tốt #tiepthilienket", s, []) == []
    flags = ai_writer.check_content("Mình đã dùng rồi, thuốc này chữa bệnh", s, [])
    assert any("từ cấm" in f for f in flags)
    assert any("Thiếu ghi chú" in f for f in flags)
    assert any("hiểu lầm" in f for f in flags)
    same = "Deal hot nồi chiên giá tốt #tiepthilienket"
    assert "Gần trùng nội dung với bài khác" in ai_writer.check_content(same, s, [same])


def test_parse_sub_ids():
    assert shopee.parse_sub_ids("p123456-b42") == ("123456", 42)
    assert shopee.parse_sub_ids("") == (None, None)


def test_score_prefers_higher_commission():
    base = {"price": 200000, "commission_rate": 0.05, "sales": 1000, "rating": 4.8}
    assert pipeline.score_product({**base, "commission_rate": 0.15}) > pipeline.score_product(base)


def test_niche_management(client):
    r = client.post("/niches", data={"name": "Đồ ăn vặt", "keywords": "khô gà, bánh tráng, khô gà",
                                     "target_pages": 5})
    assert r.status_code == 200 and "Đã tạo" in r.text
    with db.get_conn() as conn:
        assert [n["keywords"] for n in db.get_niches(conn) if n["name"] == "Đồ ăn vặt"] == [["khô gà", "bánh tráng"]]
    client.post("/niches", data={"name": "Tạm", "target_pages": 1})
    assert "Đã xoá" in client.post("/niches/delete", data={"name": "Tạm"}).text
    assert "Không xoá được" in client.post("/niches/delete", data={"name": "Mẹ & bé"}).text


def test_sync_detects_new_and_lost_pages(client):
    with db.get_conn() as conn:
        facebook.DEMO_PAGE_COUNT = 17          # tài khoản có thêm 2 page mới
        r = pipeline.sync_pages(conn)
        assert (r["total"], r["new"], r["lost"]) == (17, 2, 0)
        facebook.DEMO_PAGE_COUNT = 16          # mất quyền 1 page
        r = pipeline.sync_pages(conn)
        assert r["lost"] == 1
        assert conn.execute("SELECT status FROM pages WHERE id = '100000000016'").fetchone()[0] == "restricted"
        facebook.DEMO_PAGE_COUNT = 15


def test_assign_pages_to_niche(client):
    with db.get_conn() as conn:
        ids = [r[0] for r in conn.execute("SELECT id FROM pages WHERE niche = '' ORDER BY id")]
    assert ids, "demo phải có page chưa chọn ngành"
    assert "Mẹ &amp; bé" in client.get("/pages?niche=__none__").text or ids
    r = client.post("/pages/assign", data={"niche": "Thú cưng", "action": "assign", "ids": ids[:2]})
    assert "Đã áp ngành hàng" in r.text
    client.post(f"/pages/{ids[2]}/niche", data={"niche": "Mẹ & bé"})
    with db.get_conn() as conn:
        got = dict(conn.execute(f"SELECT id, niche FROM pages WHERE id IN (?, ?, ?)", ids[:3]).fetchall())
    assert got == {ids[0]: "Thú cưng", ids[1]: "Thú cưng", ids[2]: "Mẹ & bé"}
    r = client.post("/pages/assign", data={"niche": "Không có", "action": "assign", "ids": ids[:1]})
    assert "không tồn tại" in r.text


def _xlsx(rows):
    from io import BytesIO

    from openpyxl import Workbook
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_excel_for_niche(client):
    data = _xlsx([
        ["Tên sản phẩm", "Link sản phẩm", "Link aff", "Giá", "Mô tả", "Link ảnh", "Link video"],
        ["Bánh tráng trộn", "https://shopee.vn/Banh-trang-i.11.22", "https://s.shopee.vn/aaa", "35.000đ",
         "Vị sa tế", "https://img.example/a.jpg", "https://drive.google.com/file/d/VID1/view"],
        ["Khô gà lá chanh", "https://shopee.vn/Kho-ga-i.11.33", "https://s.shopee.vn/bbb", "89k", "", "", ""],
        ["Thiếu aff", "https://shopee.vn/X-i.11.44", "", "", "", "", ""],
        ["Video TikTok", "https://shopee.vn/Y-i.11.55", "https://s.shopee.vn/ccc", "", "", "",
         "https://www.tiktok.com/@a/video/1"],
    ])
    r = client.post("/products/import", data={"niche": "Đồ ăn vặt"},
                    files={"file": ("an-vat.xlsx", data, "application/octet-stream")})
    # Sản phẩm xác định bằng link sản phẩm; thiếu link aff vẫn lưu nhưng được cảnh báo
    assert "Đã thêm 4 sản phẩm" in r.text and "gắn 1 video" in r.text
    assert "chưa có link aff" in r.text and "TikTok" in r.text
    with db.get_conn() as conn:
        p = conn.execute("SELECT * FROM products WHERE item_id = '11.22'").fetchone()
        assert (p["name"], p["price"], p["niche"], p["aff_link"]) == ("Bánh tráng trộn", 35000, "Đồ ăn vặt",
                                                                     "https://s.shopee.vn/aaa")
        assert conn.execute("SELECT url FROM videos WHERE item_id = '11.22'").fetchone()[0].endswith("VID1/view")
        assert conn.execute("SELECT price FROM products WHERE item_id = '11.33'").fetchone()[0] == 89000
    # Nhập lại cùng file: cập nhật, không nhân đôi
    r = client.post("/products/import", data={"niche": "Đồ ăn vặt"},
                    files={"file": ("an-vat.xlsx", data, "application/octet-stream")})
    assert "Đã thêm 0 sản phẩm, cập nhật 4" in r.text
    r = client.post("/products/import", data={"niche": "Đồ ăn vặt"},
                    files={"file": ("x.csv", "Tên,Link aff\nKhông link sp,https://s.shopee.vn/q\n".encode(), "text/csv")})
    assert "thiếu link sản phẩm" in r.text


def test_import_csv_docx_and_text():
    csv_rows = importer.read_file("a.csv", "Tên,Link aff,Giá\nÁo thun,https://s.shopee.vn/x1,99000\n".encode())
    assert csv_rows == [{"name": "Áo thun", "aff_link": "https://s.shopee.vn/x1", "price": "99000"}]

    from io import BytesIO

    from docx import Document
    doc = Document()
    t = doc.add_table(rows=2, cols=3)
    for i, v in enumerate(["Tên sản phẩm", "Link aff", "Link video"]):
        t.rows[0].cells[i].text = v
    for i, v in enumerate(["Giày", "https://s.shopee.vn/g1", "https://cdn.x/g.mp4"]):
        t.rows[1].cells[i].text = v
    buf = BytesIO()
    doc.save(buf)
    assert importer.read_file("a.docx", buf.getvalue()) == [
        {"name": "Giày", "aff_link": "https://s.shopee.vn/g1", "video_url": "https://cdn.x/g.mp4"}]

    # Không có tiêu đề: đoán theo loại link
    rows = importer.read_file("a.txt", b"https://shopee.vn/Tui-i.5.6\nhttps://s.shopee.vn/t1\nhttps://cdn.x/t.mp4")
    assert rows == [{"product_link": "https://shopee.vn/Tui-i.5.6", "aff_link": "https://s.shopee.vn/t1",
                     "video_url": "https://cdn.x/t.mp4"}]
    assert shopee.name_from_link("https://shopee.vn/Tui-xach-nu-i.5.6") == "Tui xach nu"


def test_template_download(client):
    r = client.get("/products/template.xlsx?niche=Thú cưng")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    rows = importer.read_file("t.xlsx", r.content)
    assert rows[0]["aff_link"].startswith("https://s.shopee.vn/")


def test_add_product_with_uploaded_video_then_post_uses_video(client):
    r = client.post("/products/add", data={"niche": "Đồ ăn vặt", "aff_link": "https://s.shopee.vn/zzz",
                                           "product_link": "https://shopee.vn/Hat-dieu-i.77.88",
                                           "name": "Hạt điều rang muối", "price": "150k"},
                    files={"video_file": ("hat-dieu.mp4", b"fake-mp4-bytes", "video/mp4")})
    assert "Đã thêm 1 sản phẩm" in r.text and "gắn 1 video" in r.text
    item_id = importer.make_item_id({"product_link": "https://shopee.vn/Hat-dieu-i.77.88"})
    assert item_id == "77.88"
    with db.get_conn() as conn:
        v = conn.execute("SELECT * FROM videos WHERE item_id = ?", (item_id,)).fetchone()
        assert v["file_path"] and open(v["file_path"], "rb").read() == b"fake-mp4-bytes"
        # Page thuộc ngành "Đồ ăn vặt" -> bài nháp dùng video nếu sản phẩm có video, còn lại dùng ảnh
        conn.execute("UPDATE pages SET niche = 'Đồ ăn vặt', status = 'active', posts_per_day = 4 "
                     "WHERE id = (SELECT id FROM pages WHERE status = 'active' ORDER BY id LIMIT 1)")
        page_id = conn.execute("SELECT id FROM pages WHERE niche = 'Đồ ăn vặt'").fetchone()[0]
        from datetime import timedelta
        pipeline.generate_drafts(conn, db.now() + timedelta(days=3))
        posts = conn.execute("SELECT media_type, video_id, item_id, aff_link FROM posts WHERE page_id = ? "
                             "AND date(scheduled_at) = date(?)", (page_id, (db.now() + timedelta(days=3)).isoformat())
                             ).fetchall()
        media = {p["item_id"]: p["media_type"] for p in posts}
        assert media[item_id] == "video" and media["11.22"] == "video" and media["11.33"] == "photo"
        assert "11.44" not in media                       # chưa có link aff -> không tạo bài
        assert all(p["aff_link"].startswith("https://s.shopee.vn/") for p in posts)
        # Đăng bài video (demo)
        conn.execute("UPDATE posts SET status='approved', scheduled_at=? WHERE page_id=? AND item_id=?",
                     (db.now_iso(), page_id, item_id))
        pipeline.publish_due(conn)
        assert conn.execute("SELECT status FROM posts WHERE page_id=? AND item_id=?",
                            (page_id, item_id)).fetchone()[0] == "published"
    # Trang sản phẩm hiển thị video; xoá video
    assert "1 video" in client.get("/products?niche=%C4%90%E1%BB%93%20%C4%83n%20v%E1%BA%B7t").text
    client.post(f"/videos/{v['id']}/delete")


def test_direct_video_url():
    assert facebook.direct_video_url("https://drive.google.com/file/d/AbC_1-x/view?usp=sharing") == \
        "https://drive.google.com/uc?export=download&id=AbC_1-x"
    assert facebook.direct_video_url("https://cdn.site/v/clip.MP4") == "https://cdn.site/v/clip.MP4"
    assert facebook.direct_video_url("https://www.youtube.com/watch?v=x") is None


def test_cost_estimate():
    from app.services import ai_models, costs
    save = costs.estimate(80, 3, "save")
    quality = costs.estimate(80, 3, "quality")
    assert save["captions"] == 7200 and save["products"] == 720
    assert save["total_usd"] < quality["total_usd"]
    assert costs.estimate(500, 3, "save")["total_usd"] > save["total_usd"]
    haiku = costs.cost_per_caption("claude-haiku-4-5")
    assert haiku == costs.cost_per_caption("claude-haiku-4-5", batch=False) / 2
    assert haiku < costs.cost_per_caption("claude-sonnet-5") < costs.cost_per_caption("claude-opus-5")
    assert costs.cost_per_kit("claude-haiku-4-5", 512) < costs.cost_per_kit("claude-haiku-4-5", 768)
    assert ai_models.model_for("caption", "save") == "claude-haiku-4-5"
    assert ai_models.request_options("claude-haiku-4-5") == {}            # Haiku: không bật suy nghĩ
    assert ai_models.request_options("claude-sonnet-5") == {"output_config": {"effort": "low"}}


def test_batch_captions(monkeypatch):
    """Batch API: kết quả trả về không theo thứ tự, lỗi/refusal được báo đúng bài, token được ghi lại."""
    from types import SimpleNamespace as NS

    from app import config

    def message(text, stop="end_turn"):
        return NS(stop_reason=stop, model="claude-opus-5", content=[NS(type="text", text=text)],
                  usage=NS(input_tokens=600, output_tokens=900))

    class FakeBatches:
        def create(self, requests):
            self.n = len(requests)
            assert "fallbacks" not in requests[0]["params"]
            return NS(id="b1", processing_status="in_progress")

        def retrieve(self, _id):
            return NS(id="b1", processing_status="ended")

        def results(self, _id):
            return [
                NS(custom_id="2", result=NS(type="errored")),
                NS(custom_id="0", result=NS(type="succeeded", message=message("Bài số 0"))),
                NS(custom_id="1", result=NS(type="succeeded", message=message("", stop="refusal"))),
            ]

    fake = NS(messages=NS(batches=FakeBatches()))
    monkeypatch.setattr(config, "AI_ENABLED", True)
    monkeypatch.setattr(config, "AI_BATCH_MIN", 2)
    monkeypatch.setattr(ai_writer, "_client", lambda: fake)
    monkeypatch.setattr(ai_writer.time, "sleep", lambda _s: None)

    page = {"name": "P", "niche": "N", "tone": "vui"}
    product = {"name": "X", "price": 100000, "sales": 10, "rating": 5, "shop_name": "S"}
    usage = []
    out = ai_writer.write_captions([{"page": page, "product": product}] * 3, "#tag",
                                   lambda *a: usage.append(a))
    assert out[0] == ("Bài số 0\n\n#tag", None)
    assert out[1][0] is None and "từ chối" in out[1][1]
    assert out[2][0] is None and "errored" in out[2][1]
    assert usage == [("claude-opus-5", 600, 900, True), ("claude-opus-5", 600, 900, True)]


def test_facebook_login_flow(monkeypatch):
    """Đăng nhập Facebook: kiểm tra state, lưu token, nhận diện page, chặn tài khoản lạ."""
    from app import config
    monkeypatch.setattr(config, "FB_LOGIN_ENABLED", True)
    monkeypatch.setattr(config, "FB_APP_ID", "123")
    accounts = {"good": {"id": "999", "name": "Chủ page", "token": "LONG", "expires_in": 5184000},
                "other": {"id": "555", "name": "Người lạ", "token": "X", "expires_in": 5184000}}
    monkeypatch.setattr(facebook, "exchange_code", lambda code: accounts[code])
    seen_tokens = []
    monkeypatch.setattr(facebook, "list_managed_pages",
                        lambda token="": seen_tokens.append(token) or facebook.demo_pages(15))
    with db.get_conn() as conn:
        db.set_setting(conn, "owner_fb_id", "")
    with TestClient(app) as c:
        r = c.get("/auth/facebook", follow_redirects=False)
        assert r.headers["location"].startswith("https://www.facebook.com/")
        state = r.headers["location"].split("state=")[1].split("&")[0]
        r = c.get(f"/auth/facebook/callback?code=good&state=wrong", follow_redirects=False)
        assert "/login?error=" in r.headers["location"]
        c.get("/auth/facebook", follow_redirects=False)
        state = c.get("/auth/facebook", follow_redirects=False).headers["location"].split("state=")[1].split("&")[0]
        r = c.get(f"/auth/facebook/callback?code=good&state={state}")
        assert r.status_code == 200 and "nhận diện 15 page" in r.text and "Chủ page" in r.text
        assert seen_tokens[-1] == "LONG"
    with db.get_conn() as conn:
        s = db.get_settings(conn)
        assert (s["owner_fb_id"], s["owner_token"]) == ("999", "LONG")
    with TestClient(app) as c:
        state = c.get("/auth/facebook", follow_redirects=False).headers["location"].split("state=")[1].split("&")[0]
        r = c.get(f"/auth/facebook/callback?code=other&state={state}")
        assert "không có quyền" in r.text
        assert c.get("/", follow_redirects=False).status_code == 303


def test_studio_builds_images_and_video_then_posts_use_them(client):
    """Studio: ảnh gốc -> 3-5 ảnh 4:5 + video 9:16 cho mỗi phiên bản; bài đăng dùng album / video đó."""
    from PIL import Image

    from app.services import studio, video_maker
    with db.get_conn() as conn:
        item_id = conn.execute("SELECT item_id FROM products WHERE aff_link != '' AND niche = 'Mẹ & bé' "
                               "AND item_id NOT IN (SELECT item_id FROM videos) LIMIT 1").fetchone()[0]
    r = client.post(f"/studio/{item_id}/build", data={"new_brief": "1"})   # chạy nền (TestClient chờ xong)
    assert r.status_code == 200
    with db.get_conn() as conn:
        kits = conn.execute("SELECT * FROM media_kits WHERE item_id = ? ORDER BY variant", (item_id,)).fetchall()
        assert [k["status"] for k in kits] == ["ready", "ready"], [k["error"] for k in kits]
        imgs = json.loads(kits[0]["images"])
        assert 3 <= len(imgs) <= 5
        assert Image.open(imgs[0]).size == (1080, 1350)
        assert json.loads(kits[1]["images"]) != imgs                 # phiên bản 2 là file khác
        assert 8 < video_maker.probe_duration(kits[0]["video_path"]) < 16
        # Sửa chữ -> dựng lại, chữ được giữ
        client.post(f"/studio/{item_id}/brief", data={"headline": "Tiêu đề mới", "points": "Ý 1\nÝ 2\nÝ 3",
                                                      "cta": "Mua ngay", "video_lines": "A\nB\nC\nD"})
        assert studio.get_brief(conn, item_id)["headline"] == "Tiêu đề mới"

        # Tạo bài: sản phẩm có bộ media -> album hoặc video app tạo (tuỳ page)
        page_ids = [r[0] for r in conn.execute("SELECT id FROM pages WHERE niche = 'Mẹ & bé' AND status = 'active'")]
        media = {pipeline.pick_media(conn, item_id, pid, "alternate", 0)[0] for pid in page_ids}
        assert media <= {"album", "kit_video"} and media
        assert pipeline.pick_media(conn, item_id, page_ids[0], "album")[0] == "album"
        assert pipeline.pick_media(conn, item_id, page_ids[0], "video")[0] == "kit_video"
        # Đăng album (demo)
        kit_id = kits[0]["id"]
        conn.execute("""INSERT INTO posts(page_id, item_id, caption, media_type, kit_id, aff_link, status, scheduled_at,
                        created_at, updated_at) VALUES (?, ?, 'x', 'album', ?, 'https://s.shopee.vn/a', 'approved', ?, ?, ?)""",
                     (page_ids[0], item_id, kit_id, db.now_iso(), db.now_iso(), db.now_iso()))
        ok, fail = pipeline.publish_due(conn)
        assert ok >= 1
    # Trang Studio hiển thị ảnh + video; file media xem được, không lộ file ngoài thư mục
    html = client.get(f"/studio/{item_id}").text
    assert "Phiên bản 2" in html and "video.mp4" in html
    url = html.split('src="/media/m/')[1].split('"')[0]
    assert client.get("/media/m/" + url).status_code == 200
    assert client.get("/media/m/../../etc/passwd").status_code == 404
    assert client.get("/studio?status=ready").status_code == 200


def test_album_publish_calls_graph_api(monkeypatch, tmp_path):
    """Album: tải từng ảnh ở chế độ chưa đăng, rồi 1 bài gắn tất cả ảnh + bình luận link aff."""
    from app import config
    monkeypatch.setattr(config, "FB_ENABLED", True)
    calls = []

    class FakeResp:
        def __init__(self, data):
            self.status_code, self._d, self.text = 200, data, ""

        def json(self):
            return self._d

    class FakeClient:
        def __init__(self, *a, **k):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def post(self, url, data=None, files=None):
            calls.append(("POST", url.split("/", 4)[-1], dict(data or {}), bool(files)))
            if url.endswith("/photos"):
                return FakeResp({"id": f"ph{len(calls)}"})
            if url.endswith("/feed"):
                return FakeResp({"id": "page_post1"})
            return FakeResp({"id": "c1"})

        def get(self, url, params=None):
            calls.append(("GET", url.split("/", 4)[-1], dict(params or {}), False))
            return FakeResp({"permalink_url": "/p/1"})

    monkeypatch.setattr(facebook.httpx, "Client", FakeClient)
    imgs = []
    for i in range(3):
        p = tmp_path / f"{i}.jpg"
        p.write_bytes(b"jpg")
        imgs.append(str(p))
    res = facebook.publish({"id": "PAGE", "access_token": "T", "link_mode": "comment"}, "Nội dung", "https://s.shopee.vn/a",
                           images=imgs)
    assert res == {"fb_post_id": "page_post1", "permalink": "https://www.facebook.com/p/1"}
    photos = [c for c in calls if c[1].endswith("PAGE/photos")]
    assert len(photos) == 3 and all(c[2]["published"] == "false" and c[3] for c in photos)
    feed = next(c for c in calls if c[1].endswith("PAGE/feed"))
    assert feed[2]["message"] == "Nội dung" and json.loads(feed[2]["attached_media[2]"]) == {"media_fbid": "ph3"}
    comment = next(c for c in calls if c[1].endswith("page_post1/comments"))
    assert "https://s.shopee.vn/a" in comment[2]["message"]


def test_brief_uses_claude_vision_with_json_schema(monkeypatch, tmp_path):
    from types import SimpleNamespace as NS

    from PIL import Image

    from app import config
    from app.services import creative
    img = tmp_path / "a.jpg"
    Image.new("RGB", (1200, 1200), "red").save(img)
    sent = {}

    class FakeMessages:
        def create(self, **kw):
            sent.update(kw)
            text = json.dumps({"headline": "nồi chiên siêu to khổng lồ dung tích lớn cho cả nhà dùng thoải mái",
                               "subheadline": "Gọn đẹp", "points": ["ít dầu", "dễ rửa", "", "hẹn giờ", "thêm"],
                               "cta": "Mua ngay", "badge": "", "image_order": [5, 0, 0], "video_lines": ["A", "B"]})
            return NS(stop_reason="end_turn", model="claude-opus-5", content=[NS(type="text", text=text)],
                      usage=NS(input_tokens=1500, output_tokens=300))

    monkeypatch.setattr(config, "AI_ENABLED", True)
    monkeypatch.setattr(creative, "anthropic", None, raising=False)
    import anthropic
    monkeypatch.setattr(anthropic, "Anthropic", lambda **k: NS(messages=FakeMessages()))
    usage = []
    brief = creative.make_brief({"name": "Nồi", "price": 100000}, [str(img)], lambda *a: usage.append(a))
    assert sent["output_config"]["format"]["type"] == "json_schema"
    assert sent["messages"][0]["content"][1]["type"] == "image"
    assert len(brief["headline"]) <= 42 and brief["headline"].endswith("…")
    assert brief["points"] == ["Ít dầu", "Dễ rửa", "Hẹn giờ", "Thêm"]
    assert brief["image_order"] == [0]                 # chỉ số ảnh sai bị loại, không trùng
    assert usage == [("claude-opus-5", 1500, 300, False)]


def test_paste_shopee_links_detects_classifies_and_builds_media(client):
    """Dán link -> nhận diện -> tự xếp ngành (từ khoá) -> tạo ảnh + video; không đoán được ngành -> hỏi người dùng."""
    text = ("https://shopee.vn/Noi-chien-khong-dau-6L-i.555.666 | https://s.shopee.vn/aff1\n"
            "https://shopee.vn/product/1496179755/41457616922\n"
            "dòng rác không có link")
    r = client.post("/links", data={"links": text, "auto_media": "1"})
    assert "Đã nhận 2 link" in r.text
    with db.get_conn() as conn:
        jobs = {j["input"]: j for j in conn.execute("SELECT * FROM link_jobs")}
        a = jobs["https://shopee.vn/Noi-chien-khong-dau-6L-i.555.666"]
        assert a["status"] == "done", (a["step"], a["error"])
        assert a["niche"] == "Gia dụng nhà bếp" and "từ khoá" in a["step"]
        p = conn.execute("SELECT * FROM products WHERE item_id = '555.666'").fetchone()
        assert p["name"] == "Noi chien khong dau 6L" and p["aff_link"] == "https://s.shopee.vn/aff1"
        assert p["price"] > 0 and p["sales"] > 0
        assert conn.execute("SELECT COUNT(*) FROM media_kits WHERE item_id = '555.666' AND status = 'ready'"
                            ).fetchone()[0] == 2
        b = jobs["https://shopee.vn/product/1496179755/41457616922"]
        assert b["status"] == "need_niche"
    assert "Cần chọn ngành" in client.get("/products").text
    client.post(f"/links/{b['id']}/niche", data={"niche": "Thú cưng"})
    with db.get_conn() as conn:
        b = conn.execute("SELECT * FROM link_jobs WHERE id = ?", (b["id"],)).fetchone()
        assert b["status"] == "done" and b["item_id"] == "1496179755.41457616922"
        p = conn.execute("SELECT niche, aff_link FROM products WHERE item_id = ?", (b["item_id"],)).fetchone()
        assert p["niche"] == "Thú cưng" and p["aff_link"] == ""       # chưa có link aff: chưa tạo bài
    assert "Thú cưng" in client.get("/studio").text


def test_detect_product_from_short_aff_link(monkeypatch):
    """Link aff rút gọn -> theo chuyển hướng ra link sản phẩm; đọc tên/giá/ảnh/mô tả; giữ link aff."""
    from types import SimpleNamespace as NS

    from app import config
    monkeypatch.setattr(config, "SHOPEE_FETCH", True)
    monkeypatch.setattr(config, "SHOPEE_ENABLED", False)

    class FakeClient:
        def __init__(self, *a, headers=None, **k):
            self.ua = (headers or {}).get("User-Agent", "")

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def get(self, url, params=None):
            if url.startswith("https://s.shopee.vn/"):
                hop = NS(headers={"location": "https://shopee.vn/Tai-nghe-bluetooth-X1-i.77.99?sp_atk=abc"})
                return NS(history=[hop], url="https://shopee.vn/Tai-nghe-bluetooth-X1-i.77.99?sp_atk=abc")
            if "api/v4/item/get" in url:
                assert params == {"itemid": "99", "shopid": "77"}
                return NS(json=lambda: {"data": {
                    "name": "Tai nghe bluetooth X1 chống ồn", "price": 25900000000, "images": ["h1", "h2", "h3"],
                    "description": "Pin 30 giờ. Chống ồn chủ động.", "historical_sold": 5321,
                    "item_rating": {"rating_star": 4.87}, "categories": [{"display_name": "Thiết Bị Âm Thanh"}]}})
            raise AssertionError(url)

    monkeypatch.setattr(shopee.httpx, "Client", FakeClient)
    info = shopee.detect_product("https://s.shopee.vn/AbC123")
    assert info["product_link"] == "https://shopee.vn/Tai-nghe-bluetooth-X1-i.77.99"
    assert info["aff_link"] == "https://s.shopee.vn/AbC123"
    assert (info["name"], info["price"], info["sales"], info["rating"]) == ("Tai nghe bluetooth X1 chống ồn", 259000,
                                                                           5321, 4.9)
    assert info["images"] == [shopee.IMG_CDN + h for h in ("h1", "h2", "h3")]
    assert info["categories"] == ["Thiết Bị Âm Thanh"]
    with db.get_conn() as conn:
        from app.services import catalog
        assert catalog.guess_niche(conn, info["name"], info["categories"])[0] == "Âm thanh"


def test_veo_generate_calls_gemini_api(monkeypatch, tmp_path):
    """Veo: gửi ảnh sản phẩm làm khung đầu, 9:16, chờ xong, tải video về."""
    import sys
    from types import ModuleType, SimpleNamespace as NS

    from app import config
    from app.services import veo
    monkeypatch.setattr(config, "GEMINI_API_KEY", "k")
    frame = tmp_path / "f.jpg"
    frame.write_bytes(b"jpg")
    calls = {}

    class Ops:
        n = 0

        def get(self, op):
            Ops.n += 1
            return NS(done=True, error=None, result=NS(generated_videos=[NS(video="VIDEO")]), response=None)

    class Models:
        def generate_videos(self, model, source, config):
            calls.update(model=model, prompt=source.prompt, image=source.image, config=config)
            return NS(done=False)

    class Files:
        def download(self, file, destination):
            calls["download"] = (file, destination)
            open(destination, "wb").write(b"mp4")

    real_types = __import__("google.genai.types", fromlist=["types"])
    fake_genai = ModuleType("google.genai")
    fake_genai.Client = lambda api_key: NS(models=Models(), operations=Ops(), files=Files())
    fake_genai.types = real_types
    import google
    monkeypatch.setattr(google, "genai", fake_genai, raising=False)
    monkeypatch.setitem(sys.modules, "google.genai", fake_genai)
    monkeypatch.setattr(veo.time, "sleep", lambda s: None)

    out = tmp_path / "clip.mp4"
    info = veo.generate(str(frame), veo.build_prompt({"name": "Nồi chiên"}), str(out), model="veo-3.1-lite-generate-001")
    assert info == {"seconds": 8, "model": "veo-3.1-lite-generate-001", "simulated": False}
    assert calls["config"].aspect_ratio == "9:16" and calls["config"].duration_seconds == 8
    assert calls["image"].image_bytes == b"jpg" and "Nồi chiên" in calls["prompt"] and "Do not add any text" in calls["prompt"]
    assert calls["download"] == ("VIDEO", str(out)) and out.read_bytes() == b"mp4"
    assert veo.cost_per_video("veo-3.1-lite-generate-001") < veo.cost_per_video("veo-3.1-generate-preview")


def test_ai_video_in_studio_and_posts(client):
    """Studio: tạo video AI (giả lập khi chưa có key) -> bài video dùng video AI."""
    from app.services import video_maker
    with db.get_conn() as conn:
        item_id = conn.execute("SELECT item_id FROM media_kits WHERE status = 'ready' AND variant = 0").fetchone()[0]
    r = client.post(f"/studio/{item_id}/ai-video", data={"variant": "0", "style": "reveal"})
    assert r.status_code == 200
    with db.get_conn() as conn:
        kit = conn.execute("SELECT * FROM media_kits WHERE item_id = ? AND variant = 0", (item_id,)).fetchone()
        assert kit["ai_video_status"] == "ready", kit["ai_video_error"]
        assert "giả lập" in kit["ai_video_model"]
        assert 9 < video_maker.probe_duration(kit["ai_video_path"]) < 13        # clip 8s + cảnh cuối
        assert tuple(conn.execute("SELECT seconds, simulated FROM video_usage").fetchone()) == (8, 1)
        page_id = conn.execute("SELECT id FROM pages WHERE status = 'active' LIMIT 1").fetchone()[0]
        conn.execute("UPDATE media_kits SET status = 'error' WHERE item_id = ? AND variant != 0", (item_id,))
        assert pipeline.pick_media(conn, item_id, page_id, "alternate", 0)[0] == "ai_video"
        assert pipeline.pick_media(conn, item_id, page_id, "album", 0)[0] == "album"
        conn.execute("UPDATE media_kits SET status = 'ready' WHERE item_id = ?", (item_id,))
    html = client.get(f"/studio/{item_id}").text
    assert "video_ai.mp4" in html and "Video AI (Veo 3.1)" in html


def test_post_editor_in_review(client, monkeypatch):
    """Duyệt bài: sửa nội dung, đổi dạng đăng, chọn phiên bản + ảnh album, giờ đăng, link aff, AI viết lại."""
    with db.get_conn() as conn:
        kit = conn.execute("SELECT * FROM media_kits WHERE status = 'ready' ORDER BY variant DESC LIMIT 1").fetchone()
        page_id = conn.execute("SELECT id FROM pages WHERE status = 'active' LIMIT 1").fetchone()[0]
        now = db.now_iso()
        post_id = conn.execute(
            """INSERT INTO posts(page_id, item_id, caption, media_type, status, scheduled_at, aff_link, created_at,
                   updated_at) VALUES (?, ?, 'Bài cũ #tiepthilienket', 'photo', 'pending', ?, 'https://s.shopee.vn/old',
                   ?, ?)""", (page_id, kit["item_id"], now, now, now)).lastrowid
    imgs = json.loads(kit["images"])
    html = client.get(f"/review?page_id={page_id}").text
    assert f'id="post-{post_id}"' in html and "✏️ Sửa bài" in html and "Lưu & duyệt" in html
    r = client.post(f"/posts/{post_id}/edit", data={
        "caption": "Nội dung mới do mình sửa #tiepthilienket", "media_type": "album", "kit_id": str(kit["id"]),
        f"img_{kit['id']}": imgs[:2], "scheduled_at": "2030-01-02T09:30", "aff_link": "https://s.shopee.vn/new",
        "approve": "1"})
    assert "Đã lưu và duyệt bài" in r.text
    with db.get_conn() as conn:
        p = conn.execute("SELECT * FROM posts WHERE id = ?", (post_id,)).fetchone()
    assert p["caption"] == "Nội dung mới do mình sửa #tiepthilienket" and p["status"] == "approved"
    assert (p["media_type"], p["kit_id"]) == ("album", kit["id"]) and json.loads(p["custom_images"]) == imgs[:2]
    assert p["scheduled_at"].startswith("2030-01-02T09:30") and p["aff_link"] == "https://s.shopee.vn/new"

    # Ảnh không thuộc bộ media bị bỏ qua; chọn đủ ảnh = đăng cả bộ
    client.post(f"/posts/{post_id}/edit", data={"caption": p["caption"], "media_type": "album",
                                                "kit_id": str(kit["id"]), f"img_{kit['id']}": imgs + ["/etc/passwd"]})
    with db.get_conn() as conn:
        assert conn.execute("SELECT custom_images FROM posts WHERE id = ?", (post_id,)).fetchone()[0] is None
        conn.execute("UPDATE posts SET custom_images = ? WHERE id = ?", (json.dumps(imgs[:2]), post_id))
        # Đăng: album chỉ gồm ảnh đã chọn
        sent = {}
        monkeypatch.setattr(pipeline.facebook, "publish",
                            lambda page, caption, link, image_url="", video=None, images=None:
                            sent.update(images=images, link=link, caption=caption) or {"fb_post_id": "x1", "permalink": ""})
        conn.execute("UPDATE posts SET scheduled_at = ? WHERE id = ?", (db.now_iso(), post_id))
        conn.execute("UPDATE posts SET status = 'rejected' WHERE status = 'approved' AND id != ?", (post_id,))
        pipeline.publish_due(conn)
    assert sent == {"images": imgs[:2], "link": "https://s.shopee.vn/new",
                    "caption": "Nội dung mới do mình sửa #tiepthilienket"}

    # AI viết lại (demo: dùng mẫu), bài bị gắn cờ lại nếu vi phạm
    with db.get_conn() as conn:
        conn.execute("UPDATE posts SET status = 'pending' WHERE id = ?", (post_id,))
    r = client.post(f"/posts/{post_id}/rewrite", data={"preset": "shorter", "instruction": "thêm ý giao nhanh"})
    assert "AI đã viết lại bài" in r.text
    with db.get_conn() as conn:
        new_caption = conn.execute("SELECT caption FROM posts WHERE id = ?", (post_id,)).fetchone()[0]
    assert new_caption != "Nội dung mới do mình sửa #tiepthilienket" and new_caption.endswith("#tiepthilienket")
    # Bài đã đăng thì không sửa được
    with db.get_conn() as conn:
        conn.execute("UPDATE posts SET status = 'published' WHERE id = ?", (post_id,))
    client.post(f"/posts/{post_id}/edit", data={"caption": "hack"})
    with db.get_conn() as conn:
        assert conn.execute("SELECT caption FROM posts WHERE id = ?", (post_id,)).fetchone()[0] == new_caption


# ---------------- Dựng video theo nhạc ----------------

def _click_track(path, bpm=120.0, seconds=24.0, loud_from=10.0, sr=22050):
    """Nhạc giả: tiếng trống đều theo bpm, từ giây loud_from đánh to hơn (đoạn "điệp khúc")."""
    import wave

    import numpy as np
    y = np.zeros(int(sr * seconds))
    burst = np.random.default_rng(0).standard_normal(int(0.03 * sr)) * np.exp(-np.arange(int(0.03 * sr)) / sr * 120)
    beats = np.arange(0.25, seconds - 0.1, 60 / bpm)
    for b in beats:
        i = int(b * sr)
        y[i:i + len(burst)] += burst * (0.9 if b >= loud_from else 0.3)
    with wave.open(str(path), "wb") as w:
        w.setnchannels(1); w.setsampwidth(2); w.setframerate(sr)
        w.writeframes((np.clip(y, -1, 1) * 32767).astype(np.int16).tobytes())
    return beats


def _clip(path, src="testsrc2", seconds=6, size="320x240", audio=False):
    import subprocess

    from app.services import video_maker
    cmd = [video_maker.ffmpeg_exe(), "-y", "-v", "error", "-f", "lavfi", "-t", str(seconds), "-i",
           f"{src}{':' if '=' in src else '='}size={size}:rate=30"]
    if audio:
        cmd += ["-f", "lavfi", "-i", f"sine=f=440:duration={seconds}", "-shortest"]
    subprocess.run(cmd + ["-c:v", "libx264", "-preset", "ultrafast", str(path)], check=True)
    return str(path)


def test_beat_detection_and_cut_plan(tmp_path):
    import numpy as np

    from app.services import beat_editor as be
    truth = _click_track(tmp_path / "m.wav", bpm=120)
    music = be.analyze_music(str(tmp_path / "m.wav"))
    assert abs(music["bpm"] - 120) <= 2
    beats = np.array(music["beats"])
    errors = [np.min(np.abs(truth - b)) for b in beats]
    assert np.median(errors) < 0.03 and len(beats) >= len(truth) - 3      # nhịp dò được trùng nhịp thật
    start, end = be.pick_window(music, 8)
    assert start >= 9.5 and end - start == pytest.approx(8, abs=0.01)    # chọn đoạn to (điệp khúc)
    shots = be.plan_cuts(music, start, end, "fast")
    assert shots[0]["start"] == 0 and shots[-1]["end"] == pytest.approx(end - start, abs=1e-3)
    for s in shots[1:]:                                                   # mọi điểm cắt rơi đúng nhịp
        assert np.min(np.abs(beats - (s["start"] + start))) < 1e-3
    assert np.median([s["end"] - s["start"] for s in shots]) == pytest.approx(0.5, abs=0.05)
    slow = be.plan_cuts(music, start, end, "slow")
    assert len(slow) < len(shots) and all(s["end"] - s["start"] >= be.MIN_SHOT for s in slow)


def test_beat_clip_scoring_skips_dark_footage(tmp_path):
    from app.services import beat_editor as be
    good = be.analyze_clip(_clip(tmp_path / "good.mp4"))
    dark = be.analyze_clip(_clip(tmp_path / "dark.mp4", src="color=c=0x050505"))
    be.finalize_scores([good, dark])
    assert good["quality"].mean() > dark["quality"].mean() + 0.2
    shots = [{"start": i * 0.5, "end": (i + 1) * 0.5, "energy": 0.5} for i in range(6)]
    plan = be.choose_segments([good, dark], shots)
    assert all(p["clip"] == 0 for p in plan)
    starts = [p["src_start"] for p in plan]
    assert len(set(starts)) == len(starts)                                # không lặp cùng 1 đoạn


def test_beat_video_frame_exact_with_original_audio(tmp_path):
    import subprocess

    from app.services import beat_editor as be, video_maker
    _click_track(tmp_path / "m.wav", seconds=16)
    clips = [_clip(tmp_path / "a.mp4", audio=True), _clip(tmp_path / "b.mov", src="mandelbrot", size="240x320")]
    result = be.make_edit(clips, str(tmp_path / "m.wav"), str(tmp_path / "out"), seconds=5,
                          aspects=["9:16", "1:1"], pace="medium", original_audio=0.3, keep_order=True)
    out = result["outputs"]["9:16"]
    assert video_maker.probe_duration(out) == pytest.approx(5, abs=0.05)
    probe = subprocess.run([video_maker.ffmpeg_exe(), "-i", out, "-map", "0:v", "-f", "null", "-"],
                           capture_output=True, text=True).stderr
    assert "frame=  150" in probe                                        # đúng 5 giây x 30 khung, không trôi nhịp
    assert "1080x1920" in subprocess.run([video_maker.ffmpeg_exe(), "-i", out], capture_output=True, text=True).stderr
    assert video_maker.has_audio(out) and Path(result["outputs"]["1:1"]).exists()
    assert not list((tmp_path / "out").glob("_parts_*"))                 # dọn file tạm


def test_beat_page_upload_render_and_attach(client, tmp_path):
    from app.services import beat_editor as be
    _click_track(tmp_path / "m.wav", seconds=12)
    clip = _clip(tmp_path / "c.mp4", seconds=5)
    assert client.get("/beat").status_code == 200
    r = client.post("/beat", data={"seconds": "5", "aspects": ["9:16"], "pace": "auto", "flash": "1"},
                    files=[("clips", ("c.mp4", open(clip, "rb"), "video/mp4")),
                           ("music_file", ("m.wav", open(tmp_path / "m.wav", "rb"), "audio/wav"))])
    assert r.status_code == 200
    job = be.list_jobs()[0]
    assert job["status"] == "done", job["error"]
    page = client.get("/beat").text
    assert "video_9x16.mp4" in page and "Gắn vào sản phẩm" in page
    assert client.get(page.split('<video src="')[1].split("?")[0]).status_code == 200
    with db.get_conn() as conn:
        item_id = conn.execute("SELECT item_id FROM products LIMIT 1").fetchone()[0]
    client.post(f"/beat/{job['id']}/attach", data={"item_id": item_id, "aspect": "9:16"})
    with db.get_conn() as conn:
        assert conn.execute("SELECT COUNT(*) FROM videos WHERE item_id = ? AND file_path = ?",
                            (item_id, job["outputs"]["9:16"])).fetchone()[0] == 1
    client.post(f"/beat/{job['id']}/rerun", data={"seconds": "5", "aspects": ["1:1"], "pace": "slow"})
    job = be.load_job(job["id"])
    assert job["status"] == "done" and list(job["outputs"]) == ["1:1"] and job["version"] == 2
    r = client.post("/beat", files=[("clips", ("x.txt", b"no", "text/plain"))])
    assert "err=1" in str(r.url)
    client.post(f"/beat/{job['id']}/delete")
    assert be.load_job(job["id"]) is None


def test_beat_ai_rating_uses_vision(monkeypatch, tmp_path):
    from types import SimpleNamespace as NS

    import anthropic

    from app import config
    from app.services import beat_editor as be
    clip = be.analyze_clip(_clip(tmp_path / "a.mp4", seconds=4))
    sent = {}

    class FakeMessages:
        def create(self, **kw):
            sent.update(kw)
            return NS(model="claude-haiku-4-5", content=[NS(type="text", text='{"scores": [9, 1]}')],
                      usage=NS(input_tokens=900, output_tokens=20))

    monkeypatch.setattr(config, "AI_ENABLED", True)
    monkeypatch.setattr(anthropic, "Anthropic", lambda **k: NS(messages=FakeMessages()))
    usage = []
    be.ai_rate([clip], lambda *a: usage.append(a))
    images = [c for c in sent["messages"][0]["content"] if c["type"] == "image"]
    assert len(images) == 2 and sent["output_config"]["format"]["type"] == "json_schema"
    assert clip["ai"][0] > 1.0 and clip["ai"][-1] < 0.5                  # 2 giây đầu đẹp, 2 giây sau xấu
    assert usage == [("claude-haiku-4-5", 900, 20, False)]
