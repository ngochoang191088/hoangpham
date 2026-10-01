import csv

from openpyxl import load_workbook

from app import aff_hunter, viral_clips
from app.services import shopee

NOW = 1_790_000_000


def _video(vid, channel, views, hours_ago=10, likes=0, comments=0, shares=0, platform="tiktok", duration=30):
    return {"platform": platform, "id": vid, "url": f"https://www.tiktok.com/@a/video/{vid}", "group": "Mèo",
            "channel": channel, "author": channel.split("@")[-1], "title": "clip", "views": views,
            "likes": likes, "comments": comments, "shares": shares,
            "timestamp": NOW - hours_ago * 3600, "duration": duration}


# ---------------------------------------------------------------- video viral

def test_read_channels_groups_and_shortcuts(tmp_path):
    f = tmp_path / "kenh.txt"
    f.write_text("# ghi chú\n@meo1\n[Thú cưng]\nhttps://www.tiktok.com/@meo2  kênh mèo\n"
                 "https://www.douyin.com/user/MS4wABC\n# https://www.tiktok.com/@bo_qua\n"
                 "[Nấu ăn]\nhttps://www.youtube.com/@bep/shorts\n", encoding="utf-8")
    assert viral_clips.read_channels(f) == [
        ("Chung", "https://www.tiktok.com/@meo1"), ("Thú cưng", "https://www.tiktok.com/@meo2"),
        ("Thú cưng", "https://www.douyin.com/user/MS4wABC"), ("Nấu ăn", "https://www.youtube.com/@bep/shorts")]
    assert [viral_clips.platform_of(u) for _, u in viral_clips.read_channels(f)] == \
        ["tiktok", "tiktok", "douyin", "youtube"]


def test_viral_score_is_relative_to_each_channel():
    big = [_video(f"b{i}", "@big", 1_000_000) for i in range(5)] + [_video("b_hit", "@big", 3_000_000)]
    small = [_video(f"s{i}", "@small", 20_000) for i in range(5)] + [_video("s_hit", "@small", 400_000)]
    old = [_video("old", "@small", 9_000_000, hours_ago=24 * 30)]
    tiny = [_video("tiny", "@small", 500)]
    long = [_video("long", "@small", 900_000, duration=600)]
    scored = viral_clips.score_videos(big + small + old + tiny + long, NOW, days=7, min_views=10000,
                                      max_seconds=180)
    ids = [v["id"] for v in scored]
    assert ids[0] == "s_hit"                         # kênh nhỏ gấp 20 lần mức thường > kênh lớn gấp 3
    assert ids[1] == "b_hit"
    assert "old" not in ids and "tiny" not in ids and "long" not in ids
    assert scored[0]["ratio"] == 20.0


def test_engagement_breaks_ties_and_pick_limits_per_channel():
    a = _video("a", "@c", 100_000, likes=30_000, comments=2_000, shares=5_000)
    b = _video("b", "@c", 100_000, likes=1_000)
    scored = viral_clips.score_videos([a, b], NOW, 7, 1000, 0)
    assert scored[0]["id"] == "a"
    many = viral_clips.score_videos([_video(str(i), "@c", 50_000 + i) for i in range(10)]
                                    + [_video("x", "@d", 60_000)], NOW, 7, 1000, 0)
    chosen = viral_clips.pick(many, total=80, per_channel=3, done={"tiktok:9"})
    assert sum(1 for v in chosen if v["channel"] == "@c") == 3
    assert "9" not in [v["id"] for v in chosen] and "x" in [v["id"] for v in chosen]


def test_douyin_uses_likes_and_skips_photo_posts():
    db = viral_clips.DouyinBrowser.__new__(viral_clips.DouyinBrowser)
    db.download_urls = {}
    v = db._normalize({"aweme_id": "7", "desc": "mèo\nhài", "create_time": NOW, "author": {"nickname": "Mèo Béo"},
                       "statistics": {"digg_count": 5000, "comment_count": 10, "share_count": 3, "play_count": 0},
                       "video": {"duration": 15000, "play_addr": {"url_list": ["https://v/1.mp4"]}}},
                      "https://www.douyin.com/user/x", "Mèo")
    assert v["url"] == "https://www.douyin.com/video/7" and v["duration"] == 15 and v["title"] == "mèo hài"
    assert viral_clips.metric(v) == 50000 and db.download_urls["7"] == ["https://v/1.mp4"]
    assert db._normalize({"aweme_id": "8", "images": [{}]}, "u", "g") is None


def test_run_downloads_into_day_and_group_folders(tmp_path, monkeypatch):
    clips = {"https://www.tiktok.com/@meo": [_video(str(i), "https://www.tiktok.com/@meo", 20_000) for i in range(6)]
             + [_video("hit", "https://www.tiktok.com/@meo", 500_000)]}

    def fake_list(url, group, limit, cookiefile):
        return [{**v, "group": group} for v in clips[url]]

    def fake_download(video, dest, cookiefile):
        dest.write_bytes(b"mp4")
        return True

    monkeypatch.setattr(viral_clips, "list_ytdlp", fake_list)
    monkeypatch.setattr(viral_clips, "download_ytdlp", fake_download)
    out = tmp_path / "video_viral"
    chosen = viral_clips.run([("Thú cưng", "https://www.tiktok.com/@meo")], out, total=2, per_channel=5,
                             now=NOW, log=lambda *a: None)
    assert [v["id"] for v in chosen][0] == "hit" and len(chosen) == 2
    day = next(p for p in out.iterdir() if p.is_dir())
    files = sorted(p.name for p in (day / "Thú cưng").iterdir())
    assert files[0].startswith("01_tiktok_") and files[0].endswith("_hit.mp4")
    rows = list(csv.reader((day / "danh_sach.csv").open(encoding="utf-8-sig")))
    assert rows[0][0] == "Hạng" and rows[1][4].endswith("/hit") and len(rows) == 3
    assert "tiktok:hit" in (out / "da_tai.txt").read_text()

    again = viral_clips.run([("Thú cưng", "https://www.tiktok.com/@meo")], out, total=2, per_channel=5,
                            now=NOW, log=lambda *a: None)
    assert "hit" not in [v["id"] for v in again]                 # không tải lại video đã tải


def test_viral_main_creates_sample_file(tmp_path):
    f = tmp_path / "kenh.txt"
    assert viral_clips.main([str(f)]) == 1 and "[Thú cưng]" in f.read_text(encoding="utf-8")
    assert viral_clips.main([str(f)]) == 1                         # file mẫu toàn dòng ghi chú


# ---------------------------------------------------------------- săn sản phẩm aff

def _product(item_id, rate=0.12, sales=2000, rating=4.8, price=200_000, name="Hạt cho mèo"):
    return {"item_id": item_id, "name": name, "price": price, "commission_rate": rate, "sales": sales,
            "rating": rating, "shop_name": "Shop", "image_url": "", "product_link": f"https://shopee.vn/product/1/{item_id}",
            "offer_link": ""}


def test_hunt_filters_merges_and_ranks(monkeypatch):
    calls = []

    def fake_search(kw, limit, sort):
        calls.append(sort)
        if sort == shopee.SORT_COMMISSION_DESC:
            return [_product("1", rate=0.2), _product("2", rate=0.05), _product("3", sales=100),
                    _product("4", rating=4.2), _product("5", name="Thuốc trị bệnh cho mèo")]
        return [_product("1", rate=0.2), _product("6", rate=0.1, sales=20000, price=500_000)]

    monkeypatch.setattr(shopee, "search_products", fake_search)
    result = aff_hunter.hunt([("Thú cưng", "hạt cho mèo")], blacklist=["thuốc"], log=lambda *a: None)
    assert calls == [shopee.SORT_COMMISSION_DESC, shopee.SORT_SALES_DESC]
    assert [p["item_id"] for p in result] == ["6", "1"]           # bán chạy + giá cao xếp trên
    assert result[0]["group"] == "Thú cưng" and result[0]["keyword"] == "hạt cho mèo"


def test_aff_main_demo_writes_excel_and_links(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "tu_khoa.txt").write_text("[Thú cưng]\nhạt cho mèo\n[Gia dụng]\nnồi chiên\n", encoding="utf-8")
    assert aff_hunter.main(["--demo", "--top", "3"]) == 0
    xlsx = next((tmp_path / "san_pham_aff").glob("*.xlsx"))
    rows = list(load_workbook(xlsx).active.iter_rows(values_only=True))
    assert rows[0][:3] == ("Nhóm", "Mới?", "Tên sản phẩm") and len(rows) == 7 and rows[1][1] == "MỚI"
    links = next((tmp_path / "san_pham_aff").glob("*_link_aff.txt")).read_text(encoding="utf-8").split("\n")
    assert links[0] == "[Thú cưng]" and len({x for x in links if x.startswith("http")}) == 6
    assert aff_hunter.main(["--demo", "--top", "3"]) == 0
    rows = list(load_workbook(xlsx).active.iter_rows(values_only=True))
    assert not rows[1][1]                                          # lần 2: không còn là sản phẩm mới


def test_aff_main_explains_open_api_when_missing(tmp_path, capsys):
    assert aff_hunter.main([str(tmp_path / "tu_khoa.txt")]) == 1
    assert "Open API" in capsys.readouterr().out
