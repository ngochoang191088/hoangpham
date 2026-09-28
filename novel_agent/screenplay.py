"""Chế độ viết kịch bản phim: phát triển -> danh sách cảnh (duyệt) -> viết -> hội đồng chẩn đoán -> sửa."""
import json
import re
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Callable

from novel_agent import config
from novel_agent import screenplay_prompts as P
from novel_agent.memory import _write, slugify
from novel_agent.screenplay_schemas import Development, Outline, Revision, ScriptReview

SEVERITY_ORDER = {"critical": 0, "major": 1, "minor": 2}
SCENE_RE = re.compile(r"^\s*CẢNH\s+(\d+)\s*[.:]", re.MULTILINE | re.IGNORECASE)


def _block(text: str, cache: bool = False) -> dict:
    block = {"type": "text", "text": text}
    if cache:
        block["cache_control"] = {"type": "ephemeral"}
    return block


def split_scenes(text: str) -> dict[int, str]:
    """Tách văn bản kịch bản thành {số cảnh: toàn văn cảnh}."""
    marks = list(SCENE_RE.finditer(text))
    return {int(m.group(1)): text[m.start():(marks[i + 1].start() if i + 1 < len(marks) else len(text))].strip()
            for i, m in enumerate(marks)}


class ScriptProject:
    def __init__(self, path: Path):
        self.path = Path(path)
        self.state = json.loads((self.path / "project.json").read_text(encoding="utf-8"))

    @classmethod
    def create(cls, idea: str, settings: dict, dev: Development, base: Path = None) -> "ScriptProject":
        base = Path(base or config.SCRIPT_DIR)
        slug, n = slugify(dev.title), 1
        while (base / (slug if n == 1 else f"{slug}-{n}")).exists():
            n += 1
        path = base / (slug if n == 1 else f"{slug}-{n}")
        path.mkdir(parents=True)
        state = {"idea": idea, "settings": settings, "development": dev.model_dump(), "outline": None,
                 "outline_reviews": [], "scenes": {}, "reviews": [], "stage": "outline", "log": []}
        _write(path / "project.json", json.dumps(state, ensure_ascii=False, indent=2))
        return cls(path)

    def save(self) -> None:
        _write(self.path / "project.json", json.dumps(self.state, ensure_ascii=False, indent=2))
        _write(self.path / "story-bible.md", story_bible(self))

    @property
    def dev(self) -> Development:
        return Development.model_validate(self.state["development"])

    @property
    def outline(self) -> Outline | None:
        return Outline.model_validate(self.state["outline"]) if self.state["outline"] else None

    def scenes(self) -> dict[int, str]:
        return {int(k): v for k, v in sorted(self.state["scenes"].items(), key=lambda kv: int(kv[0]))}

    def script_text(self) -> str:
        return "\n\n".join(self.scenes().values())

    # ---------- hồ sơ cho AI ----------
    def dev_text(self) -> str:
        d = self.dev
        chars = "\n".join(f"- {c.name} ({c.age}) - {c.role}\n  Sinh lý: {c.physical}\n  Xã hội: {c.social}\n"
                          f"  Tâm lý: {c.psychological}\n  Muốn: {c.want}\n  Thái độ với chủ đề: {c.attitude_to_theme}\n"
                          f"  Giọng nói: {c.voice}" for c in d.characters)
        beats = "\n".join(f"- Phút {b.minute:g} - {b.name}: {b.content}" for b in d.beats)
        s = self.state["settings"]
        return (f"# HỒ SƠ PHÁT TRIỂN: {d.title}\nThể loại: {d.genre} | Thời lượng: {s['minutes']} phút\n"
                f"Logline: {d.logline}\nTiền đề: {d.premise}\nTư tưởng chủ đạo: {d.controlling_idea}\n"
                f"Lõi kịch: {d.core}\nKhát vọng / niềm tin sai lầm: {d.desire_vs_false_belief}\nVì sao lúc này: {d.why_now}\n"
                f"Luật thế giới: {d.world_rules}\n\n## Nhân vật\n{chars}\n\nThứ trói buộc: {d.binding}\n"
                f"Con dao của đối thủ: {d.opponent_weapon}\nXiềng xích của nhân vật chính: {d.protagonist_shackles}\n\n"
                f"## Cấu trúc\nKết phim: {d.ending} (kiểu kết: {d.ending_type})\nHình ảnh chủ đạo: {d.key_image}\n"
                f"Tuyến phụ: {d.subplot}\n{beats}")

    def outline_text(self) -> str:
        o = self.outline
        if not o:
            return ""
        rows = "\n".join(f"{c.number}. {c.heading} ({c.seconds}s) | {c.summary} | {c.value_open} -> {c.value_close} "
                         f"| Xung đột: {c.conflict} | {c.position}" for c in o.scenes)
        return f"# DANH SÁCH CẢNH ({len(o.scenes)} cảnh, ~{sum(c.seconds for c in o.scenes) / 60:.1f} phút)\n{rows}"


class ScriptStudio:
    def __init__(self, llm, log: Callable[[str], None] = print):
        self.llm = llm
        self.log = log

    def _beats(self, minutes: int) -> str:
        m = {f"m{p}": round(minutes * p / 100, 1) for p in (1, 5, 11, 23, 27, 50, 68, 77)}
        return P.BEAT_TABLE.format(minutes=minutes, short=P.SHORT_FILM if minutes <= 40 else "", **m)

    # ---------- Giai đoạn 1-3: tiền đề, cấu trúc, nhân vật ----------
    def create(self, idea: str, *, minutes: int = 15, genre: str = "tự chọn cho hợp ý tưởng", notes: str = "không",
               base=None) -> ScriptProject:
        self.log("🎬 Phát triển dự án: tiền đề -> kết phim -> cấu trúc -> nhân vật...")
        task = P.DEVELOP_TASK.format(idea=idea, genre=genre, minutes=minutes, notes=notes, beats=self._beats(minutes))
        dev = self.llm.parse(Development, system=P.SYSTEM, content=[_block(task)])
        project = ScriptProject.create(idea, {"minutes": minutes, "genre": genre, "notes": notes}, dev, base)
        self.log(f"   -> \"{dev.title}\": {dev.logline}")
        self.plan(project)
        return project

    # ---------- Giai đoạn 4: danh sách cảnh + biên tập cấu trúc ----------
    def plan(self, project: ScriptProject, note: str = None, rounds: int = 2) -> None:
        """Lập danh sách cảnh; có `note` thì sửa danh sách hiện tại theo yêu cầu của tác giả."""
        minutes = project.state["settings"]["minutes"]
        if note and project.outline:
            self.log("🗂️  Sửa danh sách cảnh theo yêu cầu của tác giả...")
            project.state["log"].append(f"Tác giả yêu cầu sửa danh sách cảnh: {note}")
            outline = self.llm.parse(Outline, system=P.SYSTEM, content=[
                _block(project.dev_text(), cache=True), _block(project.outline_text()),
                _block(P.OUTLINE_REVISE_TASK.format(feedback=f"- [critical] (Tác giả) {note}"))])
        else:
            self.log("🗂️  Lập danh sách cảnh...")
            outline = self.llm.parse(Outline, system=P.SYSTEM, content=[
                _block(project.dev_text(), cache=True), _block(P.OUTLINE_TASK.format(minutes=minutes, seconds=minutes * 60))])
        for round_no in range(rounds + 1):
            _renumber(outline)
            project.state["outline"] = outline.model_dump()
            project.save()
            review = self.llm.parse(ScriptReview, system=P.SYSTEM, effort=config.CRITIC_EFFORT, model=config.CRITIC_MODEL,
                                    content=[_block(project.dev_text(), cache=True), _block(project.outline_text()),
                                             _block(P.OUTLINE_REVIEW_TASK)])
            passed = _passed([review])
            project.state["outline_reviews"].append({"round": round_no, **review.model_dump()})
            self.log(f"   🔎 Biên tập cấu trúc vòng {round_no}: {review.score:g}/10 {'✅' if passed else '❌'} - {review.summary}")
            if passed or round_no == rounds:
                break
            self.log("   🛠️  Sửa danh sách cảnh theo góp ý...")
            outline = self.llm.parse(Outline, system=P.SYSTEM, content=[
                _block(project.dev_text(), cache=True), _block(project.outline_text()),
                _block(P.OUTLINE_REVISE_TASK.format(feedback=_feedback({"structure": review}, outline_mode=True)))])
        project.state["stage"] = "approve"
        project.save()

    # ---------- Giai đoạn 5: viết ----------
    def draft(self, project: ScriptProject, batch_seconds: int = 300) -> None:
        cards = project.outline.scenes
        batches, cur, acc = [], [], 0
        for c in cards:   # gom cảnh thành từng đợt ~5 phút phim để mỗi lần viết không quá dài
            if cur and acc + c.seconds > batch_seconds:
                batches.append(cur)
                cur, acc = [], 0
            cur.append(c)
            acc += c.seconds
        if cur:
            batches.append(cur)
        project.state["scenes"] = {}
        for batch in batches:
            first, last = batch[0].number, batch[-1].number
            self.log(f"✍️  Viết cảnh {first}-{last}...")
            written = project.script_text()
            text = self.llm.text(system=P.SYSTEM, content=[
                _block(project.dev_text() + "\n\n" + project.outline_text(), cache=True),
                _block(f"# KỊCH BẢN ĐÃ VIẾT (các cảnh trước)\n\n{written}" if written else "# CHƯA CÓ CẢNH NÀO"),
                _block(P.DRAFT_TASK.format(first=first, last=last, format_rules=P.FORMAT_RULES,
                                           dialogue_rules=P.DIALOGUE_RULES))])
            got = split_scenes(text)
            if not got:   # AI bỏ sót mọi dòng tiêu đề: giữ nguyên văn bản vào cảnh đầu đợt để không mất nội dung
                got = {first: f"CẢNH {first}. {batch[0].heading}\n\n{text}"}
            missing = [c.number for c in batch if c.number not in got]
            if missing:
                self.log(f"   ⚠️  Không tách được cảnh {missing}; kiểm tra lại trong file kịch bản.")
            for n, scene in got.items():
                project.state["scenes"][str(n)] = scene
            project.save()
        project.state["stage"] = "review"
        project.save()

    # ---------- Giai đoạn 6: hội đồng chẩn đoán + sửa ----------
    def _review(self, project: ScriptProject) -> dict[str, ScriptReview]:
        context = [_block(project.dev_text() + "\n\n" + project.outline_text(), cache=True),
                   _block(f"# KỊCH BẢN\n\n{project.script_text()}", cache=True)]

        def run(key):
            c = P.CRITICS[key]
            return key, self.llm.parse(ScriptReview, system=P.SYSTEM, model=config.CRITIC_MODEL, effort=config.CRITIC_EFFORT,
                                       content=[*context, _block(P.CRITIC_TASK.format(name=c["name"], checklist=c["checklist"]))])

        keys = list(P.CRITICS)
        results = [run(keys[0])]   # nhà phê bình đầu ghi cache, các nhà sau đọc cache song song
        with ThreadPoolExecutor(max_workers=len(keys)) as pool:
            results += list(pool.map(run, keys[1:]))
        return dict(results)

    def revise(self, project: ScriptProject, max_rounds: int = None) -> bool:
        max_rounds = config.MAX_REVISIONS if max_rounds is None else max_rounds
        for round_no in range(max_rounds + 1):
            reviews = self._review(project)
            avg = sum(r.score for r in reviews.values()) / len(reviews)
            passed = _passed(reviews.values())
            project.state["reviews"].append({"round": len(project.state["reviews"]), "average": round(avg, 2),
                                             "passed": passed, "reviews": {k: r.model_dump() for k, r in reviews.items()}})
            project.save()
            scores = ", ".join(f"{P.CRITICS[k]['name']} {r.score:g}" for k, r in reviews.items())
            self.log(f"🔎 Vòng {round_no}: {scores} -> TB {avg:.1f} {'✅ đạt' if passed else '❌ cần sửa'}")
            if passed or round_no == max_rounds:
                project.state["stage"] = "done" if passed else "review"
                project.save()
                return passed
            self.log("🛠️  Sửa các cảnh bị góp ý...")
            revision = self.llm.parse(Revision, system=P.SYSTEM, content=[
                _block(project.dev_text() + "\n\n" + project.outline_text(), cache=True),
                _block(f"# KỊCH BẢN\n\n{project.script_text()}"),
                _block(P.REVISE_TASK.format(feedback=_feedback(reviews), format_rules=P.FORMAT_RULES))])
            for scene in revision.scenes:
                old = project.state["scenes"].get(str(scene.number))
                if old is None:
                    continue
                text = scene.text.strip()
                if not SCENE_RE.match(text):   # AI quên dòng tiêu đề cảnh: giữ tiêu đề cũ
                    text = old.splitlines()[0] + "\n\n" + text
                project.state["scenes"][str(scene.number)] = text
            project.state["log"].append(f"Sửa vòng {round_no}: {revision.notes}")
            project.save()
            self.log(f"   Đã sửa {len(revision.scenes)} cảnh: {revision.notes}")
        return False


def _renumber(outline: Outline) -> None:
    for i, c in enumerate(outline.scenes, 1):
        c.number = i


def _passed(reviews) -> bool:
    return all(r.score >= config.PASS_SCORE and not any(i.severity == "critical" for i in r.issues) for r in reviews)


def _feedback(reviews: dict[str, ScriptReview], outline_mode: bool = False) -> str:
    issues = sorted(((i, k) for k, r in reviews.items() for i in r.issues), key=lambda x: SEVERITY_ORDER[x[0].severity])
    name = (lambda k: "Biên tập cấu trúc") if outline_mode else (lambda k: P.CRITICS[k]["name"])
    lines = [f"- [{i.severity}] ({name(k)}) Cảnh {', '.join(map(str, i.scenes)) or 'toàn bộ'} | {i.checklist_item}\n"
             f"  Vấn đề: {i.problem}\n  Cách sửa: {i.suggestion}" for i, k in issues]
    lines += [f"- Nhận xét chung ({name(k)}): {r.summary}" for k, r in reviews.items()]
    return "\n".join(lines)


# ---------- xuất file ----------
def story_bible(project: ScriptProject) -> str:
    """story-bible.md theo mẫu của skill sw-workflow: mở repo bằng Claude Code có plugin là làm tiếp được."""
    d, s = project.dev, project.state["settings"]
    stage = {"outline": "4 - danh sách cảnh", "approve": "4 - chờ tác giả duyệt danh sách cảnh",
             "review": "6 - sửa", "done": "6 - đã qua hội đồng"}[project.state["stage"]]
    nxt = {"outline": "Lập danh sách cảnh", "approve": "Tác giả đọc và sửa danh sách cảnh trong project.json rồi chạy 'script write'",
           "review": "Chạy 'script revise' hoặc sửa tay theo góp ý trong reviews", "done": "Đọc thử, quay thử"}[project.state["stage"]]
    last = project.state["reviews"][-1] if project.state["reviews"] else None
    return f"""# 《{d.title}》story bible

## Thông tin dự án
- Thể loại / thời lượng: {d.genre}, {s['minutes']} phút
- Đường vào: từ con số 0 (phim {'ngắn' if s['minutes'] <= 40 else 'dài'})
- File liên quan: project.json (dữ liệu), {project.path.name}.txt / .docx (kịch bản)

## Giai đoạn hiện tại
- Giai đoạn: {stage}
- Bước tiếp: {nxt}

## Đã chốt
- Tiền đề: {d.premise}
- Tư tưởng chủ đạo: {d.controlling_idea}
- Kết phim: {d.ending} ({d.ending_type})
- Nhân vật chính và đối thủ: {', '.join(f'{c.name} ({c.role})' for c in d.characters[:2])}; trói buộc: {d.binding}

## 1 Tiền đề và chủ đề
- Logline: {d.logline}
- Lõi kịch: {d.core}
- Khát vọng / niềm tin sai lầm: {d.desire_vs_false_belief}

## 2 Cấu trúc
- Hình ảnh chủ đạo: {d.key_image}
- Tuyến phụ: {d.subplot}
""" + "\n".join(f"- Phút {b.minute:g} {b.name}: {b.content}" for b in d.beats) + """

## 3 Nhân vật
| Nhân vật | Sinh lý | Xã hội | Tâm lý | Muốn | Thái độ với chủ đề |
|---|---|---|---|---|---|
""" + "\n".join(f"| {c.name} ({c.age}) | {c.physical} | {c.social} | {c.psychological} | {c.want} | {c.attitude_to_theme} |"
                for c in d.characters) + f"""
- Con dao của đối thủ: {d.opponent_weapon}
- Xiềng xích của nhân vật chính: {d.protagonist_shackles}

## 4 Danh sách cảnh
- {len(project.outline.scenes) if project.outline else 0} cảnh (chi tiết trong project.json)

## 5 Bản nháp
- {len(project.state['scenes'])} cảnh đã viết, ~{sum(len(t.split()) for t in project.state['scenes'].values())} chữ

## 6 Sửa
- {f"Vòng {last['round']}: TB {last['average']} ({'đạt' if last['passed'] else 'chưa đạt'})" if last else 'Chưa chấm'}

## Nhật ký quyết định
""" + "\n".join(f"- {line}" for line in project.state["log"]) + "\n"


def export(project: ScriptProject) -> list[Path]:
    d, s = project.dev, project.state["settings"]
    header = (f"{d.title.upper()}\n\nLogline: {d.logline}\nThể loại: {d.genre} - {s['minutes']} phút\n"
              f"Thể thức: đánh số cảnh kiểu Việt Nam, ~1 trang (180-220 chữ) / 1 phút phim\n\n")
    out = [project.path / f"{project.path.name}.txt"]
    out[0].write_text(header + project.script_text() + "\n", encoding="utf-8")
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        return out
    doc = Document()
    doc.add_heading(d.title, 0)
    doc.add_paragraph(f"Logline: {d.logline}")
    for scene in project.scenes().values():
        for line in scene.splitlines():
            line = line.strip()
            if not line:
                continue
            if SCENE_RE.match(line):
                doc.add_paragraph().add_run(line).bold = True
            elif re.match(r"^[A-ZÀ-Ỹ][A-ZÀ-Ỹ0-9 .'-]{0,30}(\s*\([^)]*\))?\s*:", line):
                p = doc.add_paragraph(line)
                p.paragraph_format.left_indent = Pt(72)
            else:
                doc.add_paragraph(line)
    out.append(project.path / f"{project.path.name}.docx")
    doc.save(out[1])
    return out


def status(project: ScriptProject) -> str:
    d = project.dev
    lines = [f"{d.title} ({d.genre}, {project.state['settings']['minutes']} phút) - {project.path}",
             f"  Logline: {d.logline}", f"  Giai đoạn: {project.state['stage']}"]
    if project.outline:
        lines.append(project.outline_text())
    if project.state["reviews"]:
        r = project.state["reviews"][-1]
        lines.append(f"  Hội đồng vòng {r['round']}: TB {r['average']} ({'đạt' if r['passed'] else 'chưa đạt'})")
    return "\n".join(lines)
