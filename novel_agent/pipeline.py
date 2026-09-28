"""Điều phối các agent: lập kế hoạch -> viết -> hội đồng phê bình -> sửa (lặp) -> ghi nhớ -> chương sau."""
from concurrent.futures import ThreadPoolExecutor
from typing import Callable

from novel_agent import config, prompts
from novel_agent.llm import dumps
from novel_agent.memory import Project
from novel_agent.schemas import ChapterRecord, PartPlan, Review, StoryBible

SEVERITY_ORDER = {"critical": 0, "major": 1, "minor": 2}


def _block(text: str, cache: bool = False) -> dict:
    block = {"type": "text", "text": text}
    if cache:
        block["cache_control"] = {"type": "ephemeral"}
    return block


class NovelStudio:
    def __init__(self, llm, log: Callable[[str], None] = print):
        self.llm = llm
        self.log = log

    # ---------- Kiến trúc sư ----------
    def create(self, idea: str, *, parts: int = 3, chapters_per_part: int = 8, words: int = 3000,
               genre: str = "tự chọn cho hợp ý tưởng", notes: str = "không", base=None) -> Project:
        self.log("🏛️  Kiến trúc sư đang dựng kinh thánh truyện...")
        task = prompts.BIBLE_TASK.format(idea=idea, genre=genre, parts=parts, chapters=chapters_per_part,
                                         words=words, notes=notes)
        bible = self.llm.parse(StoryBible, system=prompts.ARCHITECT_SYSTEM, content=[_block(task)])
        for i, part in enumerate(bible.parts, 1):   # đánh số lại cho chắc liên tục
            part.number = i
        settings = {"parts": parts, "chapters_per_part": chapters_per_part, "words": words, "genre": genre}
        project = Project.create(idea, settings, bible, base)
        self.log(f"   -> \"{bible.title}\": {len(bible.characters)} nhân vật, {len(bible.parts)} phần. Lưu tại {project.path}")
        return project

    def plan_part(self, project: Project, part_no: int) -> None:
        """Lập dàn ý chi tiết cho 1 phần, dựa trên những gì THỰC SỰ đã viết ở các phần trước."""
        part = next(p for p in project.bible.parts if p.number == part_no)
        count = project.state["settings"]["chapters_per_part"]
        start = project.first_chapter_of(part_no)
        self.log(f"🗺️  Kiến trúc sư lập dàn ý phần {part_no}: \"{part.title}\" (chương {start}-{start + count - 1})...")
        task = prompts.PART_PLAN_TASK.format(part_no=part_no, part_title=part.title, arc=part.arc,
                                             ending_hook=part.ending_hook, count=count, start=start,
                                             end=start + count - 1)
        plan = self.llm.parse(PartPlan, system=prompts.ARCHITECT_SYSTEM,
                              content=[_block(project.planning_context()), _block(task)])
        for i, chapter in enumerate(plan.chapters):   # đánh số lại cho chắc liên tục
            chapter.number = start + i
        project.set_part_plan(part_no, plan.chapters)

    # ---------- Nhà văn + hội đồng + sửa ----------
    def _review(self, context: list[dict], draft: str, number: int) -> dict[str, Review]:
        chapter = _block(f"# BẢN THẢO CHƯƠNG {number}\n\n{draft}", cache=True)

        def run(key):
            critic = prompts.CRITICS[key]
            task = prompts.CRITIC_TASK.format(name=critic["name"], number=number, focus=critic["focus"])
            return key, self.llm.parse(Review, system=prompts.WORKSHOP_SYSTEM, content=[*context, chapter, _block(task)],
                                       model=config.CRITIC_MODEL, effort=config.CRITIC_EFFORT)

        keys = list(prompts.CRITICS)
        results = [run(keys[0])]   # nhà phê bình đầu ghi cache ngữ cảnh, các nhà sau đọc cache song song
        with ThreadPoolExecutor(max_workers=len(keys)) as pool:
            results += list(pool.map(run, keys[1:]))
        return dict(results)

    @staticmethod
    def _passed(reviews: dict[str, Review]) -> bool:
        return all(r.score >= config.PASS_SCORE and not any(i.severity == "critical" for i in r.issues)
                   for r in reviews.values())

    @staticmethod
    def _feedback(reviews: dict[str, Review]) -> str:
        issues = sorted(((i, key) for key, r in reviews.items() for i in r.issues),
                        key=lambda x: SEVERITY_ORDER[x[0].severity])
        lines = [f"- [{i.severity}] ({prompts.CRITICS[key]['name']}) Tại: {i.location}\n  Vấn đề: {i.problem}\n"
                 f"  Cách sửa: {i.suggestion}" for i, key in issues]
        lines += [f"- Nhận xét chung của {prompts.CRITICS[k]['name']}: {r.summary}" for k, r in reviews.items()]
        return "\n".join(lines)

    def write_chapter(self, project: Project, number: int) -> dict:
        plan = project.plan_of(number)
        words = project.state["settings"]["words"]
        context = [_block(project.chapter_context(number), cache=True)]
        self.log(f"✍️  Nhà văn viết chương {number}: \"{plan.title}\"...")
        draft = self.llm.text(system=prompts.WORKSHOP_SYSTEM,
                              content=[*context, _block(prompts.WRITE_TASK.format(number=number, title=plan.title, words=words))])

        rounds, best = [], None
        for round_no in range(config.MAX_REVISIONS + 1):
            reviews = self._review(context, draft, number)
            avg = sum(r.score for r in reviews.values()) / len(reviews)
            passed = self._passed(reviews)
            rounds.append({"round": round_no, "words": len(draft.split()), "average": round(avg, 2), "passed": passed,
                           "reviews": {k: r.model_dump() for k, r in reviews.items()}})
            scores = ", ".join(f"{prompts.CRITICS[k]['name']} {r.score:g}" for k, r in reviews.items())
            self.log(f"   🔎 Vòng {round_no}: {scores} -> TB {avg:.1f} {'✅ đạt' if passed else '❌ cần sửa'}")
            if best is None or avg > best[1]:
                best = (draft, avg)
            if passed or round_no == config.MAX_REVISIONS:
                break
            self.log(f"   🛠️  Biên tập viên sửa chương {number} theo góp ý...")
            draft = self.llm.text(system=prompts.WORKSHOP_SYSTEM, content=[
                *context, _block(f"# BẢN THẢO CHƯƠNG {number}\n\n{draft}"),
                _block(prompts.REVISE_TASK.format(number=number, feedback=self._feedback(reviews), words=words))])

        final, score = (draft, rounds[-1]["average"]) if rounds[-1]["passed"] else best
        project.save_chapter(plan, final)
        project.save_reviews(number, rounds)
        self._archive(project, number, final, context)
        return {"number": number, "score": score, "passed": rounds[-1]["passed"], "rounds": len(rounds)}

    def review_only(self, project: Project, number: int) -> dict[str, Review]:
        """Cho hội đồng đọc lại một chương đã có (ví dụ sau khi bạn tự sửa tay), không thay đổi gì."""
        context = [_block(project.chapter_context(number), cache=True)]
        return self._review(context, project.chapter_text(number), number)

    # ---------- Thủ thư ----------
    def _archive(self, project: Project, number: int, text: str, context: list[dict]) -> None:
        self.log(f"📚 Thủ thư ghi nhớ chương {number} (nhân vật, sự thật, tuyến truyện)...")
        record = self.llm.parse(ChapterRecord, system=prompts.WORKSHOP_SYSTEM, content=[
            *context, _block(f"# CHƯƠNG {number} (bản cuối)\n\n{text}"), _block(prompts.RECORD_TASK.format(number=number))])
        project.apply_record(number, record)

        part_no = project.part_of(number)
        plan = project.part_plan(part_no)
        if number == plan[-1].number:   # hết phần -> nén cả phần thành 1 bản tóm tắt
            part = next(p for p in project.bible.parts if p.number == part_no)
            chapters = "\n".join(f"Chương {c.number} - {c.title}: {project.record(c.number).summary}" for c in plan)
            summary = self.llm.text(system=prompts.WORKSHOP_SYSTEM, effort=config.CRITIC_EFFORT, max_tokens=16000,
                                    content=[_block(prompts.PART_SUMMARY_TASK.format(part_no=part_no, title=part.title,
                                                                                     chapters=chapters))])
            project.state["part_summaries"][str(part_no)] = summary
            project.save()
            self.log(f"🏁 Xong phần {part_no}: \"{part.title}\"")

    # ---------- Chạy liên tục ----------
    def write(self, project: Project, count: int | None = None) -> list[dict]:
        """Viết tiếp `count` chương (None = tới hết truyện). Dừng giữa chừng thì lần sau chạy lại sẽ viết tiếp."""
        results = []
        while count is None or len(results) < count:
            nxt = project.next_chapter()
            if nxt is None:
                self.log("🎉 Đã viết xong toàn bộ tiểu thuyết.")
                break
            part_no, number = nxt
            if project.part_plan(part_no) is None:
                self.plan_part(project, part_no)
                continue
            results.append(self.write_chapter(project, number))
        return results


def status(project: Project) -> str:
    b = project.bible
    lines = [f"{b.title} ({b.genre}) - {project.path}"]
    for p in b.parts:
        plan = project.part_plan(p.number)
        if plan is None:
            lines.append(f"  Phần {p.number} \"{p.title}\": chưa lập dàn ý")
            continue
        done = [c for c in plan if project.record(c.number)]
        lines.append(f"  Phần {p.number} \"{p.title}\": {len(done)}/{len(plan)} chương")
        for c in plan:
            mark = "✅" if project.record(c.number) else "⬜"
            lines.append(f"    {mark} Chương {c.number}: {c.title}")
    lines.append(f"  Tuyến truyện đang mở: {len(project.state['open_threads'])}")
    return "\n".join(lines)


def export(project: Project) -> list:
    """Gộp các chương đã viết thành 1 file .md và 1 file .docx."""
    b = project.bible
    md = [f"# {b.title}\n"]
    doc = None
    try:
        from docx import Document
        doc = Document()
        doc.add_heading(b.title, 0)
    except ImportError:
        pass
    for p in b.parts:
        plan = [c for c in project.part_plan(p.number) or [] if project.record(c.number)]
        if not plan:
            continue
        md.append(f"\n## Phần {p.number}: {p.title}\n")
        if doc:
            doc.add_heading(f"Phần {p.number}: {p.title}", 1)
        for c in plan:
            text = project.chapter_text(c.number)
            md.append(f"\n### Chương {c.number}: {c.title}\n\n{text}\n")
            if doc:
                doc.add_heading(f"Chương {c.number}: {c.title}", 2)
                for para in text.split("\n\n"):
                    if para.strip():
                        doc.add_paragraph(para.strip())
    out = [project.path / f"{project.path.name}.md"]
    out[0].write_text("\n".join(md), encoding="utf-8")
    if doc:
        out.append(project.path / f"{project.path.name}.docx")
        doc.save(out[1])
    return out


def bible_dump(project: Project) -> str:
    return dumps(project.state["bible"])
